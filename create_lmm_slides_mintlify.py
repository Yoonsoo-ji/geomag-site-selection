# -*- coding: utf-8 -*-
"""
지자기 모델(LMM) 시범 구축(안) — Mintlify 디자인 시스템 기반 발표자료 생성.

    python create_lmm_slides.py     -> docs/output/YYYYMMDD_HHMMSS_LMM_시범구축_발표자료.pptx

디자인 근거: 2026_yoons_wiki/sources/papers/awesome-design-md/mintlify/DESIGN.md
  · 순백 캔버스 · 근검정 텍스트(#0d0d0d) · 브랜드 그린(#18E299) 절제 사용
  · 그림자 대신 5~8% 테두리로 깊이 표현 (Mintlify 는 그림자를 거의 쓰지 않는다)
  · 풀필(pill) 배지 · 16/24px 대형 라운드 카드
  · 모노 대문자 라벨 = 기술 라벨 전용 보이스, 본문은 산세리프

내용은 4-층 결합 구조(§12 CLAUDE.md)와 시범 구축 6단계 절차를 재구성한 것이다.
배경 그라디언트는 외부 자산 없이 실행 시 생성하므로 이 파일 하나로 재현된다.
"""
import math
import tempfile
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
DOCS_OUT = HERE / "docs" / "output"
ASSETS = Path(tempfile.gettempdir()) / "lmm_slide_assets"

# ── 팔레트 ────────────────────────────────────────────────────────────────
INK        = RGBColor(0x0D, 0x0D, 0x0D)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GREEN      = RGBColor(0x18, 0xE2, 0x99)
GREEN_TINT = RGBColor(0xD4, 0xFA, 0xE8)
GREEN_DEEP = RGBColor(0x0F, 0xA7, 0x6E)
G700       = RGBColor(0x33, 0x33, 0x33)
G500       = RGBColor(0x66, 0x66, 0x66)
G400       = RGBColor(0x88, 0x88, 0x88)
BORDER     = RGBColor(0xE8, 0xE8, 0xE8)   # rgba(0,0,0,.05~.08) on white
SURFACE    = RGBColor(0xFA, 0xFA, 0xFA)
D_TEXT     = RGBColor(0xED, 0xED, 0xED)
D_MUTED    = RGBColor(0xA0, 0xA0, 0xA0)
D_CARD     = RGBColor(0x16, 0x16, 0x16)
D_BORDER   = RGBColor(0x2A, 0x2A, 0x2A)

SANS = "Malgun Gothic"     # Inter 대체 (한글 포함)
MONO = "Consolas"          # Geist Mono 대체

SW, SH = 13.333, 7.5
M = 0.78                   # 좌우 여백
CW = SW - 2 * M            # 콘텐츠 폭


# ── 저수준 헬퍼 ───────────────────────────────────────────────────────────
def charspace(run, pts):
    """Inter 스타일 트래킹. 음수 = 압축(디스플레이 크기), 양수 = 확장(모노 라벨)."""
    run.font._rPr.set("spc", str(int(round(pts * 100))))


def noshadow(shape):
    """Mintlify 는 그림자를 거의 쓰지 않는다 — 테마 effectRef 까지 떼어낸다."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def has_hangul(text):
    return any(0xAC00 <= ord(c) <= 0xD7A3 or 0x3130 <= ord(c) <= 0x318F
               for c in text)


def text_width(text, size, font):
    """풀필 폭 추정 — 한글은 전각이라 라틴보다 약 두 배 넓다."""
    em = size / 72.0
    w = 0.0
    for ch in text:
        w += em * (1.02 if ord(ch) > 0x2E80 else
                   (0.62 if font == MONO else 0.55))
    return w


def rrect(slide, x, y, w, h, radius=0.20, fill=WHITE, line=BORDER, lw=1.0):
    """Mintlify 카드: 흰 배경 + 초박 테두리 + 그림자 없음."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    noshadow(sh)
    sh.text_frame.word_wrap = True
    return sh


def circle(slide, x, y, d, fill=GREEN_TINT, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.0)
    noshadow(sh)
    return sh


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size, color, *, font=SANS, bold=False, spc=0.0,
         space_before=0, space_after=0, line=None, align=None, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line is not None:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    if spc:
        charspace(r, spc)
    return p


def fill_shape_text(shape, text, size, color, *, bold=False, font=SANS,
                    spc=0.0, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    if spc:
        charspace(r, spc)


def pill(slide, x, y, text, *, size=10, fill=GREEN_TINT, color=GREEN_DEEP,
         font=MONO, bold=True, pad=0.19, h=0.30, line=None, spc=0.5):
    """풀필 배지 — Mintlify 시그니처 형태(9999px radius)."""
    if font == MONO and has_hangul(text):
        font, spc = SANS, 0.0     # Geist Mono 대체 글꼴에 한글 글리프가 없다
    w = pad * 2 + text_width(text, size, font) + abs(spc) * len(text) / 72.0
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.5
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.0)
    noshadow(sh)
    fill_shape_text(sh, text, size, color, bold=bold, font=font, spc=spc)
    return sh, w


def eyebrow(slide, x, y, text, color=GREEN_DEEP, w=6.0):
    """모노 대문자 섹션 라벨 — 콘텐츠 종류를 가르는 시각적 구분자."""
    tf = textbox(slide, x, y, w, 0.24)
    para(tf, text.upper(), 10.5, color, font=MONO, bold=True, spc=0.65, first=True)
    return tf


def build_backgrounds():
    """Mintlify 의 '대기감 그라디언트' — 녹-백 구름 워시. 실행 시 생성한다."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    W, H = 1600, 900

    def wash(path, base, target, cx, cy, rmax, power, sx=1.0, sy=1.0, blur=20):
        if Path(path).exists():
            return
        img = Image.new("RGB", (W, H), base)
        px = img.load()
        for y in range(H):
            for x in range(0, W, 2):
                d = math.hypot((x - cx) * sx, (y - cy) * sy) / rmax
                t = max(0.0, 1.0 - d) ** power
                c = tuple(int(base[k] + (target[k] - base[k]) * t)
                          for k in range(3))
                px[x, y] = c
                if x + 1 < W:
                    px[x + 1, y] = c
        img.filter(ImageFilter.GaussianBlur(blur)).save(path, "PNG")

    wash(ASSETS / "bg_light.png", (255, 255, 255), (0xC8, 0xF4, 0xDF),
         W * 0.5, H * -0.12, H * 1.15, 1.7, sy=1.05, blur=18)
    wash(ASSETS / "bg_dark.png", (0x0D, 0x0D, 0x0D), (0x10, 0x5E, 0x43),
         W * 0.72, H * 1.05, H * 1.30, 2.2, sx=0.85, blur=24)


def bg(slide, image):
    pic = slide.shapes.add_picture(str(ASSETS / image), 0, 0,
                                   Inches(SW), Inches(SH))
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footnote(slide, text, y=6.82, color=G400):
    tf = textbox(slide, M, y, CW, 0.32)
    para(tf, text, 10.5, color, line=1.35, first=True)


def slide_head(slide, eyebrow_text, title, sub=None, *, y=0.62):
    eyebrow(slide, M, y, eyebrow_text)
    tf = textbox(slide, M, y + 0.34, CW, 0.72)
    para(tf, title, 34, INK, bold=True, spc=-0.72, first=True)
    if sub:
        tf2 = textbox(slide, M, y + 1.08, CW * 0.78, 0.44)
        para(tf2, sub, 14.5, G500, line=1.45, first=True)


# ── 슬라이드 1 · 타이틀 ───────────────────────────────────────────────────
def s1_title(prs):
    s = blank(prs)
    bg(s, "bg_dark.png")

    pill(s, M, 1.62, "LOCAL MAGNETIC MODEL", size=10.5,
         fill=D_CARD, color=GREEN, line=D_BORDER, h=0.34, pad=0.24)

    tf = textbox(s, M, 2.28, CW * 0.86, 1.9)
    para(tf, "지자기 모델(LMM)", 54, D_TEXT, bold=True, spc=-1.28,
         line=1.12, first=True)
    para(tf, "시범 구축 절차(안)", 54, GREEN, bold=True, spc=-1.28, line=1.12)

    tf2 = textbox(s, M, 4.62, CW * 0.62, 1.0)
    para(tf2, "전지구 기준장·지역장·지각 이상·외부장을 물리적 성분별로 분리하고\n"
              "다시 결합해, 단순 내삽으로는 닿지 않는 정밀도에 도달한다.",
         15, D_MUTED, line=1.55, first=True)

    # 4개 층 이름을 모노 풀필로 — 이후 슬라이드까지 이어지는 시각 모티프
    x = M
    for name in ("IGRF", "REGIONAL", "CRUSTAL", "EXTERNAL"):
        _, w = pill(s, x, 5.92, name, size=10, fill=D_CARD, color=D_MUTED,
                    line=D_BORDER, h=0.32, pad=0.22)
        x += w + 0.14

    tf3 = textbox(s, SW - M - 3.0, 6.02, 3.0, 0.3, align=PP_ALIGN.RIGHT)
    para(tf3, "2026. 07", 10.5, G500, font=MONO, spc=0.6, first=True)

    s.notes_slide.notes_text_frame.text = (
        "LMM 은 하나의 통짜 모델이 아니라 파장이 다른 네 개의 물리 성분을 "
        "따로 계산해 합치는 구조다. 이 발표는 그 구조와 시범 구축 절차를 다룬다.")
    return s


# ── 슬라이드 2 · 구성 방정식 ──────────────────────────────────────────────
def s2_equation(prs):
    s = blank(prs)
    bg(s, "bg_light.png")

    eyebrow(s, M, 0.72, "THE MODEL")
    tf = textbox(s, M, 1.06, CW, 0.8)
    para(tf, "네 개 층의 합으로 자기장을 기술한다", 34, INK, bold=True,
         spc=-0.72, first=True)

    card = rrect(s, M, 2.14, CW, 1.44, radius=0.32, fill=WHITE)
    tfc = card.text_frame
    tfc.margin_left = tfc.margin_right = Inches(0.3)
    tfc.margin_top = tfc.margin_bottom = 0
    tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tfc.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    subs = {"LMM", "IGRF", "Regional", "Crustal", "External"}
    for txt in ("B", "LMM", "(r, t)   =   ", "B", "IGRF", "   +   ",
                "B", "Regional", "   +   ", "B", "Crustal", "   +   ",
                "B", "External"):
        r = p.add_run()
        r.text = txt
        r.font.name = SANS
        r.font.color.rgb = GREEN_DEEP if txt in subs else INK
        if txt in subs:
            r.font.size = Pt(14)
            r.font.bold = False
            r.font._rPr.set("baseline", "-22000")   # 진짜 아래첨자
        else:
            r.font.size = Pt(27)
            r.font.bold = True
            charspace(r, -0.5)

    items = [
        ("파장", "층마다 담당하는 공간 파장이 다르다 — 수천 km 부터 수십 m 까지"),
        ("시간", "기준장은 5년 주기, 외부장은 1분 단위로 변한다"),
        ("독립성", "성분을 분리해야 어느 층의 자료를 보강할지 판단할 수 있다"),
    ]
    cw = (CW - 2 * 0.34) / 3
    for i, (label, desc) in enumerate(items):
        x = M + i * (cw + 0.34)
        rrect(s, x, 3.94, cw, 1.86, radius=0.22, fill=WHITE)
        tfx = textbox(s, x + 0.30, 4.28, cw - 0.60, 1.20)
        para(tfx, label, 14, INK, bold=True, spc=-0.2, first=True)
        para(tfx, desc, 11.5, G500, line=1.55, space_before=7)

    footnote(s, "※ 각 층 자료를 물리적 성분별로 분리·결합하여 단순 내삽의 한계를 보완한다.",
             y=6.34)
    s.notes_slide.notes_text_frame.text = (
        "핵심은 덧셈식 자체가 아니라 '왜 나누는가'다. 파장·시간 스케일이 다르고, "
        "분리해 두어야 부족한 층을 골라 보강할 수 있다.")
    return s


# ── 슬라이드 3 · 층별 구성 ────────────────────────────────────────────────
def s3_layers(prs):
    s = blank(prs)
    slide_head(s, "LAYERS · 01–04", "층별 구성과 자료원",
               "네 개 층은 서로 다른 자료원에서 오고, 각자 다른 파장 대역을 책임진다.")

    layers = [
        ("01", "Core", "IGRF-14", "전지구 기준장",
         "5년 주기 갱신 · degree ≤ 13\n지구 외핵이 만드는 주자기장"),
        ("02", "Regional", "지표 절대측정", "17점 반복관측",
         "50 ~ 3,000 km 파장 보정\n반복관측점 기반 지역 편차"),
        ("03", "Crustal", "KIGAM 항공자력", "자력이상도",
         "0.05 ~ 50 km 지각 이상\n암상·지질구조가 만드는 국소 이상"),
        ("04", "External", "상시관측 · 청양(CYG)", "1분 자료",
         "외부장 · 시간변화 분리\n전리층·자기권 기원 일변화"),
    ]
    cw = (CW - 0.40) / 2
    ch = 2.05
    for i, (num, en, src, tag, desc) in enumerate(layers):
        x = M + (i % 2) * (cw + 0.40)
        y = 2.28 + (i // 2) * (ch + 0.34)
        rrect(s, x, y, cw, ch, radius=0.24, fill=WHITE)

        circle(s, x + 0.30, y + 0.30, 0.46)
        tfn = textbox(s, x + 0.30, y + 0.30, 0.46, 0.46,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        para(tfn, num, 12, GREEN_DEEP, font=MONO, bold=True, first=True)

        tft = textbox(s, x + 0.90, y + 0.30, cw - 1.2, 0.34)
        para(tft, en, 19, INK, bold=True, spc=-0.2, first=True)

        pill(s, x + 0.90, y + 0.72, src, size=9.5, fill=SURFACE, color=G500,
             line=BORDER, h=0.28, pad=0.17)

        tfd = textbox(s, x + 0.30, y + 1.22, cw - 0.60, 0.72)
        para(tfd, desc, 11.5, G700, line=1.55, first=True)

        # 태그는 카드 우상단에 모노로 — 자료 규모를 한눈에
        tfg = textbox(s, x + cw - 2.3 - 0.30, y + 0.34, 2.3, 0.26,
                      align=PP_ALIGN.RIGHT)
        para(tfg, tag, 10, G400, font=MONO, spc=0.4, first=True)

    s.notes_slide.notes_text_frame.text = (
        "01 은 전지구 모델을 그대로 쓰고, 02~04 가 국내 자료로 채워지는 부분이다. "
        "현재 02 는 17점, 04 는 청양 한 곳에 의존한다.")
    return s


# ── 슬라이드 4 · 구축 절차 ────────────────────────────────────────────────
def s4_process(prs):
    s = blank(prs)
    slide_head(s, "PROCESS", "시범 구축 절차(안)",
               "자료 수집에서 산출물까지 여섯 단계로 진행한다.")

    steps = [
        ("01", "입력자료 수집",
         "IGRF-14 · 지표 절대측정 17점\nKIGAM 항공자력 · 청양(CYG) 1분"),
        ("02", "층별 성분 산출",
         "기준장 · 지역 · 지각 · 시간변화\n네 개 성분으로 분리"),
        ("03", "결합 및 잔차면 구성",
         "4개 층 결합 후 관측점 잔차 내삽\n반복관측 기반 공간분포 보정"),
        ("04", "최적계산",
         "편각 D · 복각 I · 총자력 F 산출\n2025 기준시점으로 환산"),
        ("05", "교차검증 (LOO)",
         "관측점을 하나씩 빼며 재적합 (Leave-One-Out)\nKPI 점검 후 채택 여부 판단"),
        ("06", "산출물",
         "도엽별 자침편차\n웹 · 엑셀 계산기"),
    ]
    cw = (CW - 2 * 0.34) / 3
    ch = 1.96
    for i, (num, title, desc) in enumerate(steps):
        x = M + (i % 3) * (cw + 0.34)
        y = 2.24 + (i // 3) * (ch + 0.32)
        last = (i == 5)
        rrect(s, x, y, cw, ch, radius=0.22,
              fill=GREEN_TINT if last else WHITE,
              line=GREEN_TINT if last else BORDER)

        pill(s, x + 0.26, y + 0.26, num, size=10, h=0.30, pad=0.16,
             fill=WHITE if last else GREEN_TINT,
             color=GREEN_DEEP, line=None)

        tft = textbox(s, x + 0.26, y + 0.70, cw - 0.52, 0.40)
        para(tft, title, 14, INK, bold=True, spc=-0.2, line=1.25, first=True)

        tfd = textbox(s, x + 0.26, y + 1.20, cw - 0.52, 0.62)
        para(tfd, desc, 10.5, GREEN_DEEP if last else G500, line=1.5, first=True)

    footnote(s, "※ 각 단계는 앞 단계 산출물을 소비한다 — 자료가 보강되면 02 부터 다시 흐른다.",
             y=6.80)
    s.notes_slide.notes_text_frame.text = (
        "절차는 한 번 돌고 끝나는 파이프라인이 아니라 순환이다. "
        "05 에서 KPI 미달이면 01 의 자료를 늘려 다시 돌린다.")
    return s


# ── 슬라이드 5 · 최적계산과 교차검증 (KPI) ────────────────────────────────
def s5_kpi(prs):
    s = blank(prs)
    slide_head(s, "STEP 04 · 05", "최적계산과 교차검증")

    lw = CW * 0.46
    tf = textbox(s, M, 2.20, lw, 2.6)
    para(tf, "결합된 잔차면에서 편각 D · 복각 I · 총자력 F 를 산출하고, "
             "2025 기준시점으로 환산한다.", 14, G700, line=1.65, first=True)
    para(tf, "검증은 Leave-One-Out 방식이다. 관측점을 하나씩 제외하고 "
             "나머지로 모델을 다시 적합한 뒤, 빠진 점을 얼마나 맞히는지 본다. "
             "17점처럼 표본이 적을 때 과적합을 걸러내는 표준적인 방법이다.",
         14, G700, line=1.65, space_before=14)

    x = M
    for name in ("편각 D", "복각 I", "총자력 F"):
        _, w = pill(s, x, 5.24, name, size=10.5, font=SANS, spc=0,
                    fill=SURFACE, color=G500, line=BORDER, h=0.32, pad=0.20)
        x += w + 0.14

    # KPI 대형 수치 콜아웃
    rx = M + CW * 0.52
    rw = CW * 0.48
    kpis = [
        ("DECLINATION D", "< 0.1°", "도폭 단위 자침편차를 단일 값으로 표기하기 위한 상한"),
        ("TOTAL FIELD F", "< 50 nT", "총자력 산출값이 실측을 대신할 수 있는 허용 오차"),
    ]
    for i, (label, value, note) in enumerate(kpis):
        y = 2.20 + i * 2.02
        rrect(s, rx, y, rw, 1.80, radius=0.28, fill=WHITE)
        tfl = textbox(s, rx + 0.36, y + 0.28, rw - 0.72, 0.26)
        para(tfl, label, 10, G400, font=MONO, bold=True, spc=0.65, first=True)

        tfv = textbox(s, rx + 0.36, y + 0.62, rw - 0.72, 0.66)
        para(tfv, value, 38, GREEN_DEEP, bold=True, spc=-1.0, line=1.0,
             first=True)

        tfn = textbox(s, rx + 0.36, y + 1.36, rw - 0.72, 0.30)
        para(tfn, note, 10.5, G500, line=1.4, first=True)

    footnote(s, "※ KPI 는 공학적 목표치다. 미달 시 모델 구조가 아니라 "
                "입력자료의 밀도와 시각 정보를 먼저 의심한다.", y=6.52)
    s.notes_slide.notes_text_frame.text = (
        "두 KPI 중 편각이 훨씬 까다롭다. 외부장 일변화만으로도 0.1도를 넘길 수 있어 "
        "관측 시각 기록이 없으면 원리적으로 달성이 어렵다.")
    return s


# ── 슬라이드 6 · 산출물 ───────────────────────────────────────────────────
def s6_outputs(prs):
    s = blank(prs)
    slide_head(s, "STEP 06 · DELIVERABLES", "산출물",
               "모델은 두 가지 형태로 사용자에게 도달한다.")

    outs = [
        ("도엽별 자침편차",
         "지형도 도폭 단위로 자침편차를 산출한다. 도폭 안에서 편각 변화가 "
         "표기 정밀도보다 작아야 단일 값 표기가 성립하므로, 채택 축척은 "
         "모델 오차 예산에 따라 결정된다.",
         ["도폭 단위 단일 값 산출", "지형도 난외 표기 자료로 활용"]),
        ("웹 · 엑셀 계산기",
         "좌표와 표고를 넣으면 편각 D · 복각 I · 총자력 F 를 즉시 돌려준다. "
         "웹은 외부 의존 없는 단일 파일, 엑셀은 수식이 열려 있어 검증이 가능하다.",
         ["오프라인 단일 파일 웹 계산기", "수식 공개형 스프레드시트"]),
    ]
    cw = (CW - 0.44) / 2
    for i, (title, body, bullets) in enumerate(outs):
        x = M + i * (cw + 0.44)
        rrect(s, x, 2.30, cw, 3.62, radius=0.34, fill=WHITE)

        circle(s, x + 0.38, 2.66, 0.52)
        tfn = textbox(s, x + 0.38, 2.66, 0.52, 0.52,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        para(tfn, "0" + str(i + 1), 12, GREEN_DEEP, font=MONO, bold=True, first=True)

        tft = textbox(s, x + 0.38, 3.38, cw - 0.76, 0.40)
        para(tft, title, 21, INK, bold=True, spc=-0.24, first=True)

        tfb = textbox(s, x + 0.38, 3.90, cw - 0.76, 1.10)
        para(tfb, body, 12, G500, line=1.62, first=True)

        for j, b in enumerate(bullets):
            y = 5.06 + j * 0.38
            circle(s, x + 0.40, y + 0.10, 0.11, fill=GREEN)
            tfl = textbox(s, x + 0.66, y, cw - 1.04, 0.30)
            para(tfl, b, 11.5, G700, first=True)

    footnote(s, "※ 두 산출물 모두 동일한 모델 파일을 소비한다 — 재적합하면 함께 갱신된다.",
             y=6.28)
    s.notes_slide.notes_text_frame.text = (
        "산출물이 모델 파일을 참조하도록 만들어 두면 수치를 하드코딩하지 않아도 되고, "
        "재적합 결과가 자동으로 반영된다.")
    return s


# ── 슬라이드 7 · 마무리 ───────────────────────────────────────────────────
def s7_close(prs):
    s = blank(prs)
    bg(s, "bg_dark.png")

    pill(s, M, 1.52, "WHY IT MATTERS", size=10.5, fill=D_CARD, color=GREEN,
         line=D_BORDER, h=0.34, pad=0.24)

    tf = textbox(s, M, 2.18, CW * 0.82, 1.9)
    para(tf, "각 층 자료를 물리적 성분별로 분리·결합하여", 32, D_TEXT,
         bold=True, spc=-0.7, line=1.28, first=True)
    para(tf, "단순 내삽의 한계를 보완한다", 32, GREEN, bold=True, spc=-0.7, line=1.28)

    cards = [
        ("분리", "파장·시간 스케일이 다른 성분을 섞지 않는다"),
        ("결합", "관측점 잔차면으로 층 사이 빈틈을 메운다"),
        ("검증", "Leave-One-Out 으로 과적합 없이 성능을 읽는다"),
    ]
    cw = (CW - 2 * 0.34) / 3
    for i, (label, desc) in enumerate(cards):
        x = M + i * (cw + 0.34)
        rrect(s, x, 4.42, cw, 1.42, radius=0.22, fill=D_CARD, line=D_BORDER)
        tfx = textbox(s, x + 0.30, 4.72, cw - 0.60, 0.92)
        para(tfx, label, 15, GREEN, bold=True, spc=-0.2, first=True)
        para(tfx, desc, 11, D_MUTED, line=1.5, space_before=6)

    tfe = textbox(s, M, 6.24, CW, 0.3)
    para(tfe, "지자기 모델(LMM) 시범 구축(안)  ·  ", 10.5, G500, spc=0.2, first=True)
    r = tfe.paragraphs[0].add_run()
    r.text = "2026. 07"
    r.font.size = Pt(10.5)
    r.font.name = MONO
    r.font.color.rgb = G500
    charspace(r, 0.6)

    s.notes_slide.notes_text_frame.text = (
        "마무리 메시지는 하나다 — 층 분리는 복잡함을 위한 복잡함이 아니라, "
        "어디를 보강해야 정밀도가 오르는지 알기 위한 구조다.")
    return s


def main():
    build_backgrounds()

    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    s1_title(prs)
    s2_equation(prs)
    s3_layers(prs)
    s4_process(prs)
    s5_kpi(prs)
    s6_outputs(prs)
    s7_close(prs)

    DOCS_OUT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = DOCS_OUT / f"{stamp}_LMM_시범구축_발표자료.pptx"
    prs.save(out)
    print("saved:", out)
    return out


if __name__ == "__main__":
    main()
