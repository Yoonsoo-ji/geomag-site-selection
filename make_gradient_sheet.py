# -*- coding: utf-8 -*-
r"""GSM-19T 원문을 붙여 넣으면 수평·수직 구배가 나오는 시트 (2026-09-21)

    python make_gradient_sheet.py --into "docs/output/…_3판_전체50.xlsx"

발주자가 만든 `지자기 자기구배계산_입력출력.xlsx` 는 **지점마다 12줄씩 골라 붙여넣는**
구조였다. 실제 기기 파일은 5초 간격 26~42회를 한 블록으로 쏟아 내므로(10초·12줄 가정과
다르다) 사람이 잘라 붙이는 일이 끝나지 않는다. 그래서 **파일 내용을 통째로 붙여넣고
블록을 시트가 알아서 자르도록** 다시 짰다.

## 읽는 규칙 (2026-09-21 발주자 확정)

| 항목 | 규칙 |
|---|---|
| 블록 나누기 | `/time` 줄마다 새 블록. 그 앞 머리글 줄(`/Gem …`·`/ID …`)은 잘려 있어도 된다 |
| 거리 배정 | **블록을 시작 시각순으로** 줄 세워 P0(전) → 1 m → 2 m → … → P0(후) |
| 최대 거리 | 블록 수에서 자동으로 나온다(블록 n개 → n−2 m 까지). 오호는 동 6 · 서 10 · 남 4 · 북 8 m |
| 빼는 읽음 | F = 0.00 · 품질 뒷자리 0 · **그 지점 중앙값에서 ±5 nT 벗어난 값** |
| 시각 | 그 지점 **첫 읽음 시각**(hhmm) — 카드 시각 칸 서식과 같다 |
| 수직 | 삼각대 파일의 블록을 시각순으로 **봉 4 → 3 → 2 → 1** |

⚠️ **서쪽 파일은 블록이 파일 안에서 11 → 10 → 12 … 로 섞여 있었다.** 파일 순서가 아니라
**시각순**이 맞다(발주자 확인). 그래서 시작 시각으로 순위를 매겨 거리를 배정한다 —
파일이 어떤 순서로 덤프됐든 결과가 같다.

⚠️ **MINIFS·MAXIFS·TEXTSPLIT 를 쓰지 않는다.** 현장 PC 의 엑셀 판을 모르므로
`SMALL/LARGE(IF(…))` 배열식과 `SUBSTITUTE(…,REPT(" ",30))` 토막내기처럼 오래된 판에서도
도는 식만 쓴다. 배열식은 openpyxl `ArrayFormula` 로 넣어 CSE 로 저장된다.

⚠️ 결과는 **카드와 같은 모양**(12행 × 8열 · 수직 4행)으로 내므로 그대로 복사해
카드 `B16` · `C42` 에 **값으로** 붙여 넣으면 된다. 카드 칸을 수식으로 묶지 않는 이유는,
지점을 바꿔 다음 자료를 붙여 넣는 순간 앞 지점 값이 사라지기 때문이다.
"""
from __future__ import annotations

import argparse
import datetime as dt
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

import make_trial_field_card as MC
import trial_survey_points as TP

ROOT = Path(__file__).parent
OUT_DIR = ROOT / "docs" / "output"
SHEET = "⑨ 구배 자동계산"

DIRS = ["동", "서", "남", "북"]
DATASETS = DIRS + ["수직"]
PASTE_TOP, PASTE_N = 9, 700          # 붙여넣기 A9:E708 (블록 12개 × 40줄 + 머리글 여유)
SUM_TOP, SUM_N = 9, 14               # 블록 요약 14줄 — 한 방향 최대 블록 수(P0+10 m+P0 = 12)
RES_TOP = 9                          # 결과 표 G9:O20 (P0(전)·1~10 m·P0(후))
VERT_TOP = 23                        # 수직 결과 G23:K26 (봉 4·3·2·1) — 21 제목 · 22 머리
PRE_TOP = 29                         # 미리보기 G29:N33
SITE_COL = 91                        # CM — 지점 드롭다운 목록
SEG_COL = 95                         # CQ 구간 이름표 · CR~CU 방향별 구간구배
VERT_BONG = [(4, 185), (3, 140), (2, 95), (1, 50)]
TRIM_NT = 5                          # 중앙값에서 이만큼 벗어나면 뺀다

HEAD = PatternFill("solid", fgColor="16324F")
SUBHEAD = PatternFill("solid", fgColor="D9E2F3")
PASTE_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="E2EFDA")
F_TITLE = Font(name=MC.FONT, size=14, bold=True, color="FFFFFF")
F_H = Font(name=MC.FONT, size=10, bold=True, color="1F3864")
F_SM = Font(name=MC.FONT, size=9, color="44546A")
F_CALC = Font(name=MC.FONT, size=10, bold=True, color="1F3864")
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _cols(d: int) -> dict:
    """자료 하나(방향 또는 수직)가 쓰는 숨긴 열 — 줄마다 다섯 칸."""
    base = 17 + 5 * d                                    # Q 부터
    return dict(zip(("blk", "t", "f", "q", "use"), (CL(base + i) for i in range(5))))


def _sum_cols(d: int) -> dict:
    """블록 요약 아홉 칸."""
    base = 43 + 9 * d                                    # AQ 부터
    keys = ("k", "n", "t0", "t1", "med", "used", "avg", "drop", "rank")
    return dict(zip(keys, (CL(base + i) for i in range(len(keys)))))


def build(wb: Workbook, sites: list[str]) -> None:
    ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = "1D4ED8"
    ws.sheet_view.showGridLines = False
    for c, w in [("A", 30), ("B", 30), ("C", 30), ("D", 30), ("E", 30), ("F", 2),
                 ("G", 13), ("H", 11), ("I", 13), ("J", 11), ("K", 13), ("L", 12),
                 ("M", 13), ("N", 11), ("O", 24)]:
        ws.column_dimensions[c].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1         # 결과·수직·미리보기가 가로 한 쪽에 들어간다
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.merge_cells("A1:O1")
    ws["A1"].value = "⑨ 구배 자동계산 — GSM-19T 파일 내용을 그대로 붙여 넣는 곳"
    ws["A1"].font, ws["A1"].fill, ws["A1"].alignment = F_TITLE, HEAD, AL_L
    ws.row_dimensions[1].height = 26

    guide = (
        "① 아래 노란 칸(A9·B9·C9·D9·E9)에 그 방향 txt 파일 내용을 «통째로» 붙여 넣습니다"
        "(메모장에서 Ctrl+A → Ctrl+C → 여기서 그 칸 하나만 고르고 Ctrl+V).   "
        "② 오른쪽 초록 표가 바로 채워집니다 — 블록을 시각순으로 줄 세워 P0(전)·1 m·2 m…·P0(후) 로 "
        "배정하고, 0.00 읽음과 품질 뒷자리 0, 중앙값에서 ±5 nT 벗어난 값을 빼고 평균을 냅니다.   "
        "③ 결과 표 H9:O20 을 복사해 그 지점 카드의 B16 에 «값으로 붙여넣기», 수직 I23:I26 은 카드 C42 에 "
        "같은 방법으로 붙여 넣습니다.   ④ 다음 지점을 하려면 노란 칸을 모두 지우고 다시 붙여 넣으세요."
    )
    ws.merge_cells("A2:O2")
    ws["A2"].value, ws["A2"].font, ws["A2"].alignment = guide, F_SM, AL_L
    ws.row_dimensions[2].height = 58

    ws["A4"].value, ws["A4"].font = "지점", F_H
    ws["B4"].fill, ws["B4"].font, ws["B4"].alignment = PASTE_FILL, F_CALC, AL_C
    ws["B4"].border = MC.BOX
    dv = DataValidation(type="list", formula1="=$CM$9:$CM$58", allowBlank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    for i, name in enumerate(sites):
        ws.cell(9 + i, SITE_COL).value = name            # CM 열 — 드롭다운 목록
    ws.column_dimensions[CL(SITE_COL)].hidden = True
    ws.merge_cells("C4:O4")
    ws["C4"].value = ('=IF($B$4="","붙여 넣기 전에 지점을 골라 주세요 — 결과를 어느 카드에 옮길지 '
                      '알려 드립니다","결과 표 H9:O20 → 「"&$B$4&"」 시트의 B16 에 값으로 붙여넣기  ·  '
                      '수직 I23:I26 → 같은 시트 C42")')
    ws["C4"].font, ws["C4"].alignment = F_CALC, AL_L

    # ── 붙여넣기 칸 ────────────────────────────────────────────
    for d, name in enumerate(DATASETS):
        col = CL(1 + d)
        c = ws[f"{col}8"]
        c.value = f"{name} 원문 붙여넣기" if d < 4 else "수직(삼각대) 원문 붙여넣기"
        c.font, c.fill, c.alignment, c.border = Font(name=MC.FONT, size=10, bold=True,
                                                     color="FFFFFF"), HEAD, AL_C, MC.BOX
        for r in range(PASTE_TOP, PASTE_TOP + 3):        # 맨 위 세 줄만 노랗게 — 어디 붙일지 보이게
            ws[f"{col}{r}"].fill = PASTE_FILL
    ws.row_dimensions[8].height = 22
    ws.freeze_panes = "A9"

    # ── 줄마다 푸는 숨긴 열 ────────────────────────────────────
    for d in range(len(DATASETS)):
        src, h = CL(1 + d), _cols(d)
        for r in range(PASTE_TOP, PASTE_TOP + PASTE_N):
            s = f"{src}{r}"
            tok = f'SUBSTITUTE(TRIM({s})," ",REPT(" ",30))'
            prev = f'N({h["blk"]}{r - 1})' if r > PASTE_TOP else "0"
            ws[f'{h["blk"]}{r}'] = (f'=IF({s}="",{prev},'
                                    f'IF(LEFT(TRIM({s}),5)="/time",{prev}+1,{prev}))')
            ws[f'{h["t"]}{r}'] = (f'=IFERROR(IF(ISNUMBER(--LEFT(TRIM({s}),1)),'
                                  f'VALUE(LEFT(TRIM({s}),6)),""),"")')
            ws[f'{h["f"]}{r}'] = f'=IFERROR(VALUE(TRIM(MID({tok},31,30))),"")'
            ws[f'{h["q"]}{r}'] = f'=IFERROR(TRIM(MID({tok},61,30)),"")'
            # ⚠️ 품질 칸이 비면 RIGHT("",1)="" 이라 «0 이 아니므로» 통과해 버린다
            ws[f'{h["use"]}{r}'] = (f'=IF(OR({h["t"]}{r}="",{h["f"]}{r}="",{h["q"]}{r}=""),0,'
                                    f'IF(AND({h["f"]}{r}>1000,RIGHT({h["q"]}{r},1)<>"0"),1,0))')
        for k in h.values():
            ws.column_dimensions[k].hidden = True

    # ── 블록 요약 ──────────────────────────────────────────────
    lo, hi = PASTE_TOP, PASTE_TOP + PASTE_N - 1
    for d in range(len(DATASETS)):
        h, s = _cols(d), _sum_cols(d)
        B = f"${h['blk']}${lo}:${h['blk']}${hi}"
        T = f"${h['t']}${lo}:${h['t']}${hi}"
        F = f"${h['f']}${lo}:${h['f']}${hi}"
        U = f"${h['use']}${lo}:${h['use']}${hi}"
        for i in range(SUM_N):
            r = SUM_TOP + i
            k = f"{s['k']}{r}"
            ws[k] = i + 1
            ws[f"{s['n']}{r}"] = f'=COUNTIFS({B},{k},{U},1)'
            # ⚠️ else 를 비워 두면 FALSE 가 0 으로 섞일 수 있다 — "" 로 둬야 무시된다
            cond = f'IF(({B}={k})*({U}=1),{F},"")'
            # ⚠️ 시각은 «유효 판독»이 아니라 «그 블록의 모든 판독»에서 잡는다.
            #    유효 판독으로 잡으면 한 블록이 통째로 탈락했을 때 t0 가 비고,
            #    RANK 가 그 블록을 건너뛰어 «뒤 거리가 한 칸씩 당겨진다»(코덱스 C1).
            tcond = f'IF({B}={k},{T},"")'
            ws[f"{s['t0']}{r}"] = ArrayFormula(f"{s['t0']}{r}",
                                               f'=IFERROR(SMALL({tcond},1),"")')
            ws[f"{s['t1']}{r}"] = ArrayFormula(f"{s['t1']}{r}",
                                               f'=IFERROR(LARGE({tcond},1),"")')
            ws[f"{s['med']}{r}"] = ArrayFormula(f"{s['med']}{r}",
                                                f'=IFERROR(MEDIAN({cond}),"")')
            med = f"{s['med']}{r}"
            rng = (f'{B},{k},{U},1,{F},">="&{med}-{TRIM_NT},{F},"<="&{med}+{TRIM_NT}')
            ws[f"{s['used']}{r}"] = f'=IF({med}="",0,COUNTIFS({rng}))'
            ws[f"{s['avg']}{r}"] = (f'=IF({s["used"]}{r}=0,"",'
                                    f'ROUND(AVERAGEIFS({F},{rng}),2))')
            ws[f"{s['drop']}{r}"] = (f'=IF({s["med"]}{r}="","",'
                                     f'COUNTIFS({B},{k},{T},">0")-{s["used"]}{r})')
            # ⚠️ RANK.EQ 를 쓰면 #NAME? 이 된다 — 2010 이후 함수는 파일에 `_xlfn.` 을 달고
            #    저장돼야 해서, 엑셀 밖에서 쓴 이름을 엑셀이 못 알아본다. 옛 이름 RANK 를 쓴다.
            ws[f"{s['rank']}{r}"] = (f'=IF({s["t0"]}{r}="","",'
                                     f'RANK({s["t0"]}{r},${s["t0"]}${SUM_TOP}:'
                                     f'${s["t0"]}${SUM_TOP + SUM_N - 1},1))')
        for key in s.values():
            ws.column_dimensions[key].hidden = True

    # ── 결과 표 (카드와 같은 모양) ─────────────────────────────
    heads = ["거리"] + [f"{x} {y}" for x in DIRS for y in ("시각", "평균 F(nT)")]
    for j, t in enumerate(heads):
        c = ws.cell(8, 7 + j)
        c.value, c.font, c.fill, c.alignment, c.border = t, F_H, SUBHEAD, AL_C, MC.BOX
    labels = ["P0(전)"] + [f"{m} m" for m in range(1, 11)] + ["P0(후)"]
    for i, lab in enumerate(labels):
        r = RES_TOP + i
        c = ws.cell(r, 7)
        c.value, c.font, c.alignment, c.border = lab, F_H, AL_C, MC.BOX
        ws.row_dimensions[r].height = 19
        for d in range(4):
            s = _sum_cols(d)
            nb = f'COUNT(${s["t0"]}${SUM_TOP}:${s["t0"]}${SUM_TOP + SUM_N - 1})'
            order = "1" if i == 0 else (nb if lab == "P0(후)" else str(i + 1))
            guard = ("" if i == 0 else
                     (f'IF({nb}<2,"",' if lab == "P0(후)" else f'IF({nb}-1<{i + 1},"",'))
            close = "" if guard == "" else ")"
            rk = f'${s["rank"]}${SUM_TOP}:${s["rank"]}${SUM_TOP + SUM_N - 1}'
            pick = lambda col: (  # noqa: E731 — 순서로 블록을 찾아 값을 꺼낸다
                f'IFERROR(INDEX(${col}${SUM_TOP}:${col}${SUM_TOP + SUM_N - 1},'
                f'MATCH({order},{rk},0)),"")')
            tcell = ws.cell(r, 8 + 2 * d)
            fcell = ws.cell(r, 9 + 2 * d)
            tcell.value = f'={guard}IFERROR(INT({pick(s["t0"])}/100),""){close}'
            fcell.value = f'={guard}{pick(s["avg"])}{close}'
            tcell.number_format = '00":"00'
            fcell.number_format = "#,##0.00"
            for c in (tcell, fcell):
                c.font, c.fill, c.alignment, c.border = F_CALC, CALC_FILL, AL_C, MC.BOX

    # ── 수직 결과 ──────────────────────────────────────────────
    ws.merge_cells(f"G{VERT_TOP - 2}:O{VERT_TOP - 2}")
    ws[f"G{VERT_TOP - 2}"].value = "수직 — 삼각대 파일 블록을 시각순으로 봉 4 → 3 → 2 → 1"
    ws[f"G{VERT_TOP - 2}"].font, ws[f"G{VERT_TOP - 2}"].fill = F_H, SUBHEAD
    ws[f"G{VERT_TOP - 2}"].alignment = AL_L
    # ⚠️ 「읽은 시각」·「비고」는 한 칸에 안 들어간다 — 칸을 합쳐 쓴다(결과 표 열 폭은 그대로)
    for j, t in enumerate(["봉 번호", "높이", "평균 F(nT)"]):
        c = ws.cell(VERT_TOP - 1, 7 + j)
        c.value, c.font, c.fill, c.alignment, c.border = t, F_H, SUBHEAD, AL_C, MC.BOX
    for span, t in (("J", "읽은 시각"), ("L", "비고")):
        rng = f"{span}{VERT_TOP - 1}:" + ("K" if span == "J" else "N") + f"{VERT_TOP - 1}"
        ws.merge_cells(rng)
        c = ws[f"{span}{VERT_TOP - 1}"]
        c.value, c.font, c.fill, c.alignment = t, F_H, SUBHEAD, AL_C
        MC._span(ws, rng, None)
    s = _sum_cols(4)
    for i, (bong, cm) in enumerate(VERT_BONG):
        r = VERT_TOP + i
        ws.cell(r, 7).value = f"봉 {bong}"
        ws.cell(r, 8).value = f"{cm} cm"
        rk = f'${s["rank"]}${SUM_TOP}:${s["rank"]}${SUM_TOP + SUM_N - 1}'
        idx = lambda col: (  # noqa: E731
            f'IFERROR(INDEX(${col}${SUM_TOP}:${col}${SUM_TOP + SUM_N - 1},'
            f'MATCH({i + 1},{rk},0)),"")')
        ws.cell(r, 9).value = f'={idx(s["avg"])}'
        ws.cell(r, 9).number_format = "#,##0.00"
        # ⚠️ TEXT(…,"00\:00\:00") 는 #VALUE! 가 됐다 — escape 를 쓰지 말고 토막내 붙인다
        hms = lambda v: (  # noqa: E731 — hhmmss 숫자를 11:16:32 로
            f'TEXT(INT({v}/10000),"00")&":"&TEXT(MOD(INT({v}/100),100),"00")'
            f'&":"&TEXT(MOD({v},100),"00")')
        ws.cell(r, 10).value = (f'=IF({idx(s["t0"])}="","",{hms(idx(s["t0"]))}'
                                f'&"~"&{hms(idx(s["t1"]))})')
        ws.cell(r, 12).value = (f'=IF({idx(s["used"])}=0,"",{idx(s["used"])}&"회 평균"'
                                f'&IF({idx(s["drop"])}>0," · "&{idx(s["drop"])}&"회 제외","")'
                                f'&IF({idx(s["used"])}<10," · 읽음이 적습니다 — 재측정 확인",""))')
        ws.merge_cells(f"J{r}:K{r}")
        ws.merge_cells(f"L{r}:N{r}")
        for cc in range(7, 15):
            c = ws.cell(r, cc)
            c.border, c.alignment = MC.BOX, AL_C
            c.font = F_CALC if cc == 9 else MC.F_VAL
            if cc == 9:
                c.fill = CALC_FILL
        ws.cell(r, 12).alignment = AL_L
        ws.row_dimensions[r].height = 30      # 비고가 두 줄이 된다(제외·재측정 안내)
    note_r = VERT_TOP + len(VERT_BONG)
    ws.merge_cells(f"G{note_r}:O{note_r}")
    ws[f"G{note_r}"].value = ("높이는 카드 기본값(185·140·95·50 cm)입니다. 실제로 다르면 카드에서 "
                              "고쳐 적어 주세요. 「중심점 1.0 m·1.5 m」처럼 따로 잰 파일은 여기 넣지 "
                              "말고 카드 비고에 적습니다.")
    ws[f"G{note_r}"].font, ws[f"G{note_r}"].alignment = F_SM, AL_L
    ws.row_dimensions[note_r].height = 17

    # ── 미리보기 (카드에 붙이기 전에 눈으로 확인) ───────────────
    ws.merge_cells(f"G{PRE_TOP - 1}:O{PRE_TOP - 1}")
    ws[f"G{PRE_TOP - 1}"].value = "붙여넣기 전 확인 — 카드에서 나올 값과 같습니다"
    ws[f"G{PRE_TOP - 1}"].font, ws[f"G{PRE_TOP - 1}"].fill = F_H, SUBHEAD
    ws[f"G{PRE_TOP - 1}"].alignment = AL_L
    for j, t in enumerate(["방향", "잰 거리", "P0 전후차", "F 변화폭", "최대 구간구배", "발생구간"]):
        c = ws.cell(PRE_TOP, 7 + j)
        c.value, c.font, c.fill, c.alignment, c.border = t, F_H, SUBHEAD, AL_C, MC.BOX
    ws.merge_cells(f"M{PRE_TOP}:N{PRE_TOP}")
    for ref, t in ((f"M{PRE_TOP}", "읽음/제외"), (f"O{PRE_TOP}", "확인")):
        c = ws[ref]
        c.value, c.font, c.fill, c.alignment = t, F_H, SUBHEAD, AL_C
    MC._span(ws, f"M{PRE_TOP}:N{PRE_TOP}", None)
    ws[f"O{PRE_TOP}"].border = MC.BOX
    for i in range(10):
        ws.cell(SUM_TOP + i, SEG_COL).value = ("P0~1 m" if i == 0 else f"{i}~{i + 1} m")
    ws.column_dimensions[CL(SEG_COL)].hidden = True
    for d in range(4):
        r = PRE_TOP + 1 + d
        s = _sum_cols(d)
        fc = CL(9 + 2 * d)
        seg = CL(SEG_COL + 1 + d)                         # CR~CU — 방향별 구간구배
        for i in range(10):
            rr = SUM_TOP + i
            a, b = f"${fc}${RES_TOP + i}", f"${fc}${RES_TOP + i + 1}"
            ws[f"{seg}{rr}"] = f'=IF(OR({a}="",{b}=""),"",ABS({b}-{a}))'
        ws.column_dimensions[seg].hidden = True
        rng = f"${seg}${SUM_TOP}:${seg}${SUM_TOP + 9}"
        lab = f"${CL(SEG_COL)}${SUM_TOP}:${CL(SEG_COL)}${SUM_TOP + 9}"
        nb = f'COUNT(${s["t0"]}${SUM_TOP}:${s["t0"]}${SUM_TOP + SUM_N - 1})'
        hh = _cols(d)
        B0 = f'${hh["blk"]}${PASTE_TOP}:${hh["blk"]}${PASTE_TOP + PASTE_N - 1}'
        T0 = f'${hh["t"]}${PASTE_TOP}:${hh["t"]}${PASTE_TOP + PASTE_N - 1}'
        ws.cell(r, 7).value = DIRS[d]
        ws.cell(r, 8).value = f'=IF({nb}<3,"",({nb}-2)&" m")'
        ws.cell(r, 9).value = (f'=IF(COUNT(${fc}${RES_TOP},${fc}${RES_TOP + 11})<2,"",'
                               f'${fc}${RES_TOP + 11}-${fc}${RES_TOP})')
        # ⚠️ 값이 하나면 MAX−MIN 이 0 이 되어 «변화가 없었다»로 읽힌다 — 두 개부터
        ws.cell(r, 10).value = (f'=IF(COUNT(${fc}${RES_TOP}:${fc}${RES_TOP + 10})<2,"",'
                                f'MAX(${fc}${RES_TOP}:${fc}${RES_TOP + 10})'
                                f'-MIN(${fc}${RES_TOP}:${fc}${RES_TOP + 10}))')
        ws.cell(r, 11).value = f'=IF(COUNT({rng})=0,"",MAX({rng}))'
        ws.cell(r, 12).value = (f'=IF({CL(11)}{r}="","",INDEX({lab},MATCH({CL(11)}{r},{rng},0)))')
        ws.cell(r, 13).value = (f'=IF({nb}=0,"",SUM(${s["used"]}${SUM_TOP}:'
                                f'${s["used"]}${SUM_TOP + SUM_N - 1})&" / "&'
                                f'SUM(${s["drop"]}${SUM_TOP}:${s["drop"]}${SUM_TOP + SUM_N - 1}))')
        # ⚠️ 방향 전체 합으로 보면 나쁜 지점 하나가 묻힌다 — 지점(블록)마다 센다
        used_r = f'${s["used"]}${SUM_TOP}:${s["used"]}${SUM_TOP + SUM_N - 1}'
        drop_r = f'${s["drop"]}${SUM_TOP}:${s["drop"]}${SUM_TOP + SUM_N - 1}'
        thin = f'COUNTIFS({used_r},">0",{used_r},"<10")'
        noisy = f'SUMPRODUCT(({drop_r}<>"")*({used_r}>0)*({drop_r}>{used_r}))'
        ws.merge_cells(f"M{r}:N{r}")
        orphan = f'COUNTIFS({B0},0,{T0},">0")'
        ws.cell(r, 15).value = (f'=IF({orphan}>0,"첫 /time 이 없습니다 — 머리글째 다시",'
                                f'IF({nb}=0,"원문을 붙여 넣어 주세요",'
                                f'IF({nb}>12,"블록 "&{nb}&"개 — 10 m 까지만 표시",'
                                f'IF({nb}<3,"블록이 "&{nb}&"개뿐 — P0 전·후가 다 있는지 보세요",'
                                f'IF({thin}>0,{thin}&"개 지점 읽음 부족 — 확인",'
                                f'IF({noisy}>0,{noisy}&"개 지점 제외 과다 — 확인","정상"))))))')
        for cc in range(7, 16):
            c = ws.cell(r, cc)
            c.border, c.alignment, c.font = MC.BOX, AL_C, MC.F_VAL
        ws.cell(r, 9).number_format = "#,##0.00"
        ws.cell(r, 10).number_format = "#,##0.00"
        ws.cell(r, 11).number_format = "#,##0.00"
        ws.row_dimensions[r].height = 19

    ws.merge_cells(f"G{PRE_TOP + 6}:O{PRE_TOP + 8}")
    ws[f"G{PRE_TOP + 6}"].value = (
        "· 거리 배정은 블록을 «시작 시각순»으로 줄 세워 P0(전) → 1 m → … → P0(후) 로 합니다. "
        "파일 안 블록 순서가 뒤섞여 있어도(오호 서쪽이 그랬습니다) 시각으로 다시 세웁니다.\n"
        "· 한 지점에서 뺀 읽음은 F=0.00, 품질 뒷자리 0, 그 지점 중앙값에서 ±5 nT 벗어난 값입니다. "
        "뺀 개수는 「읽음/제외」에 나옵니다.\n"
        "· 여기 값은 시간변화를 빼기 전 값이라 P0 전후차가 크면 그만큼 섞여 있다는 뜻입니다. "
        "국내 기준은 아직 없으니 값이 크다고 현장에서 자리를 접지 마세요.")
    ws[f"G{PRE_TOP + 6}"].font, ws[f"G{PRE_TOP + 6}"].alignment = F_SM, AL_L
    ws.print_area = f"G8:O{PRE_TOP + 9}"


def inject(target: Path, out: Path, tmp: Path) -> None:
    """엑셀이 시트를 옮겨 붙인다 — 웹 쿼리·그래프를 지키기 위해서다."""
    import pythoncom
    import win32com.client as win32

    shutil.copy(target, out)
    pythoncom.CoInitialize()
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.AskToUpdateLinks = False
    try:
        dst = xl.Workbooks.Open(str(out), UpdateLinks=0)
        src = xl.Workbooks.Open(str(tmp), UpdateLinks=0, ReadOnly=True)
        for i in range(dst.Worksheets.Count, 0, -1):
            if dst.Worksheets(i).Name == SHEET:
                dst.Worksheets(i).Delete()
        # ⚠️ `Copy(After=…)` 는 win32com 에서 조용히 «새 통합문서»로 새어 나간다(빈 Sheet1 만 남는다).
        #    `Before=` 로 넣을 것 — 총괄 다음 자리 = 첫 카드 시트 앞.
        after = [i for i, s in enumerate(dst.Worksheets, 1) if s.Name == "총괄"][0]
        src.Worksheets(SHEET).Copy(Before=dst.Worksheets(after + 1))
        src.Close(False)
        dst.Worksheets(SHEET).Activate()
        dst.Worksheets(SHEET).Range("A9").Select()
        dst.Save()
        dst.Close(True)
    finally:
        xl.Quit()
        pythoncom.CoUninitialize()


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--into", required=True)
    ap.add_argument("--tag", default="구배자동계산")
    a = ap.parse_args()

    pts = TP.load_points()
    sites = [f"{i:02d}_{MC._safe(p['지점명'])}" for i, p in enumerate(pts, 1)]
    wb = Workbook()
    wb.remove(wb.active)
    build(wb, sites)
    tmp = Path(tempfile.mkdtemp()) / "gradient.xlsx"
    wb.save(tmp)

    target = Path(a.into).resolve()
    out = OUT_DIR / f"{dt.datetime.now():%Y%m%d_%H%M%S}_시험탐사_현장야장_3판_{a.tag}.xlsx"
    inject(target, out, tmp)
    print(f"[저장] {out}  ({out.stat().st_size / 1e6:.2f} MB)")
    return out


if __name__ == "__main__":
    main()
