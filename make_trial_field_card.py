# -*- coding: utf-8 -*-
"""
시험 탐사 **현장용** 야장 카드 — `docs/output/*_시험탐사_현장야장.xlsx`
=========================================================================

    python make_trial_field_card.py

## 왜 따로 만드는가

`make_trial_survey_book.py` 가 내는 사무실용 야장은 **자료 관리 스키마**다 —
관측행 결합키·검증 게이트·검토 추적까지 담아 사후 분석이 깨지지 않게 한다.
그러나 현장 작업자에게는 과하다. 2,600행짜리 표를 들고 산에 오를 수는 없다.

이 파일은 **현장에서 실제로 적는 것만** 남긴 A4 가로 1~2쪽짜리 카드다.
점마다 시트 하나이며, 이미 103장을 채워 본 「현장조사 카드」와 같은 어법을 쓴다.

| 남긴 것 (현장) | 뺀 것 (사무실) |
|---|---|
| 도착·기상·Kp·장비번호 | 관측행 결합키(Observation ID 등) |
| 기준점 GNSS 실측 좌표 | 좌표계·측지기준·정확도 메타 |
| 수평 자기구배 4방향 격자 | 시간변화 선형보간 계산 |
| 수직 자기구배 높이별 | Set 검증 게이트 |
| 방위표지 정·반 시준 | 참방위각 결정방법·이전값 대비 차 |
| **현장 사진 9장** (중심점·측정모습·4방위·표지2·교란요소) | §13② 이격 증빙 4행/점 |
| 중심점 최종 위치·사유 · 현장 소견 | 기존점 성분별 대조가능성 · 검토 반영 추적표 |

⚠️ **절대측정(D·I)은 시험 탐사에 포함되지 않는다**(2026-09-07 사용자 확인).
원 계획 31면의 선점실시 절차가 「GNSS 수신기 → 오버하우저 자기구배 확인 →
중심점 선정 및 표지 설치」이며 DI-flux 절대관측은 여기 없다 — 그것은 선점이
끝난 뒤의 «관측» 단계다. 따라서 §19①(1일 6회)·§20(정수차 30′·20분)도 이
단계의 요건이 아니다. 이 야장이 잰 것은 **오버하우저 총자력뿐**이다.

⚠️ **시간변화 보정 계산은 현장에서 하지 않는다.** 현장은 «시각과 값»을 정확히
남기는 데 집중하고, P0 선형보간·구배 산출·판정은 사무실에서 17시트 야장으로
한다. 현장에는 판단에 필요한 최소 지표(P0 전후차 · 관측값 max−min)만 자동
계산해 둔다.

⚠️ 규격은 `trial_survey_spec.py` 단일 출처다 — 원 계획(중간보고 자료)에서 옮긴 값.

## Codex Delta Review 반영 (2026-09-07, Critical 3 · Major 7)

코덱스가 **실제 파일을 열어** 본 검토라 지적이 구체적이었다. 반영한 것:

| 지적 | 반영 |
|---|---|
| **C1** 카드는 F 한 칸인데 사무실은 F1·F2·F3 을 요구 — 점당 192회 산정이 성립 안 함 | **1 판독 기준으로 통일**했다. 오버하우저는 기기 내부에서 평균하므로 IAGA 예시도 점당 단일값이다. 판독이 의심스러우면 「재측」 칸에 적고 사유를 남긴다 — 작업량도 **64회/점**으로 다시 셌다 |
| **C2** Visit·Location·Set·Attempt 식별자가 없어 재시작·재방문·중심 이동을 유일하게 구분 못 함 | 카드 머리에 **Visit ID · Location ID**, 구역마다 **Set ID · 차수**. 중심점을 옮기면 **새 Location ID 로 새 카드**를 쓰고 기존 값을 덮지 않는다 |
| **C3** §18③ 자기오염 방지(작업자 금속물 제거)가 없다 | 「0. 측정 전 확인」 구역 신설 — 개인 금속물·차량·전자기기 통제 체크. **이건 D·I 와 무관하게 총자력 측정에 직접 영향을 준다** |
| **M1** 절대측정 제외의 잔재(야간 의무 문구) | 「야간」 의무를 **「시간변화가 안정된 구간」**으로 바꿨다 |
| **M3** 자동지표 넷만으로는 재측정 여부를 못 정함 | **현장 완전성 점검** 줄 신설 — 필수행·시각·P0 구간·Variometer 연결을 채웠는지 스스로 확인 |
| **M4** 방위표지 구역이 폐합차만 보고 참방위각을 재현 못 함 | 표지 ID·좌표·GNSS 취득방법·정확도·기존 확정값·결정방법·변경 사유 추가 |
| **M5** 파일명만으로 사진 대응이 보장되지 않음 | **명명규칙** `{Site}_{Visit}_{코드}_{연번}` 안내 + 촬영시각 칸 |
| **M7** 여러 쪽으로 나뉘면 2쪽 이후 지점 식별이 사라짐 | 1~3행을 **모든 쪽에 반복**(`print_title_rows`) · 머리글에 Site ID · 꼬리말에 쪽번호 |

⚠️ **판독 횟수는 발주자 확인 사항으로 남는다** — 1회(기기 내부평균)로 갈지
3회 수기로 갈지에 따라 작업량이 64 ↔ 192 회로 3배 달라진다.
"""
from __future__ import annotations

import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import trial_survey_points as TP
import trial_survey_spec as SP

ROOT = Path(__file__).parent
OUT_DIR = ROOT / "docs" / "output"

FONT = "맑은 고딕"
F_TITLE = Font(name=FONT, size=14, bold=True, color="FFFFFF")
F_SUB = Font(name=FONT, size=9, color="FFFFFF")
F_SEC = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_LBL = Font(name=FONT, size=9.5, bold=True)
F_VAL = Font(name=FONT, size=10)
F_SM = Font(name=FONT, size=8.5, color="555555")
F_WARN = Font(name=FONT, size=8.5, bold=True, color="B2182B")
F_CALC = Font(name=FONT, size=10, bold=True, color="1F3864")

FILL_TITLE = PatternFill("solid", fgColor="1F3864")
FILL_SEC = PatternFill("solid", fgColor="4472C4")
FILL_LBL = PatternFill("solid", fgColor="F2F2F2")
FILL_FIELD = PatternFill("solid", fgColor="FFF8E1")
FILL_CALC = PatternFill("solid", fgColor="E8F0E8")
FILL_P0 = PatternFill("solid", fgColor="DDEBF7")
FILL_WARN = PatternFill("solid", fgColor="FDE9E7")

AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_TL = Alignment(horizontal="left", vertical="top", wrap_text=True)
_thin = Side(style="thin", color="AAAAAA")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

COLS = "ABCDEFGHI"          # 9열 — A4 가로 1쪽
WIDTHS = [13, 13, 12, 13, 12, 13, 12, 13, 12]


def _s(ws, r, text, note=None):
    """구역 머리."""
    ws.merge_cells(f"A{r}:I{r}")
    c = ws[f"A{r}"]
    c.value, c.font, c.fill, c.alignment = text, F_SEC, FILL_SEC, AL_L
    ws.row_dimensions[r].height = 20
    if note:
        ws.merge_cells(f"A{r+1}:I{r+1}")
        n = ws[f"A{r+1}"]
        n.value, n.font, n.alignment = note, F_SM, AL_TL
        ws.row_dimensions[r + 1].height = 13 * max(1, -(-len(note) // 108)) + 3
        return r + 2
    return r + 1


def _lbl(ws, ref, t):
    if ":" in ref:
        ws.merge_cells(ref)
        ref = ref.split(":")[0]
    c = ws[ref]
    c.value, c.font, c.fill, c.alignment, c.border = t, F_LBL, FILL_LBL, AL_C, BOX


def _fld(ws, ref, fmt=None, calc=None):
    if ":" in ref:
        ws.merge_cells(ref)
        ref = ref.split(":")[0]
    c = ws[ref]
    c.border, c.alignment = BOX, AL_C
    if calc:
        c.value, c.font, c.fill = calc, F_CALC, FILL_CALC
    else:
        c.font, c.fill = F_VAL, FILL_FIELD
    if fmt:
        c.number_format = fmt
    return c


def _dv(ws, ref, items):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"',
                       allow_blank=True, showDropDown=False)
    ws.add_data_validation(d)
    d.add(ref)


def _photo(ws, r, span, hint, rows=8):
    """사진 붙일 자리. 조사자 입력 영역이므로 기입란과 같은 배경색을 쓴다."""
    a, b = span
    ws.merge_cells(f"{a}{r}:{b}{r+rows-1}")
    c = ws[f"{a}{r}"]
    c.value, c.font, c.fill, c.alignment = (
        hint, Font(name=FONT, size=9.5, color="999999"), FILL_FIELD, AL_C)
    for rr in range(r, r + rows):
        for cc in COLS[COLS.index(a):COLS.index(b) + 1]:
            ws[f"{cc}{rr}"].border = BOX
            ws[f"{cc}{rr}"].fill = FILL_FIELD
        ws.row_dimensions[rr].height = 15


def _safe(n):
    return re.sub(r'[\\/*?:\[\]]', "", str(n))[:26]


# ══════════════════════════════════════════════════════════════
def build_card(wb, p, idx):
    ws = wb.create_sheet(f"{idx:02d}_{_safe(p['지점명'])}")
    for col, w in zip(COLS, WIDTHS):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0          # 세로는 자연스럽게 넘긴다
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # ⚠️ 여러 쪽으로 나뉘면 2쪽부터 어느 지점인지 사라진다 — 머리 3행을 반복한다.
    ws.print_title_rows = "1:3"
    ws.oddHeader.right.text = f"{p['_sid']} {p['지점명']}"
    ws.oddHeader.right.size = 9
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 9

    # ── 머리 ────────────────────────────────────────────────
    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = f"[{p['_sid']}] {p['지점명']} — 지자기 시험 탐사 현장야장"
    c.font, c.fill, c.alignment = F_TITLE, FILL_TITLE, AL_L
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:I2")
    c = ws["A2"]
    caution = SP.SITE_CAUTION.get(p["지점명"], "")
    c.value = (f"구분 {p['구분']} · 원번호 {p['번호']} · 도엽 {p['도엽번호']} "
               f"{p['도엽명']} · 사전좌표 {p['위도']:.6f} / {p['경도']:.6f} · "
               f"표고 {p['표고'] if p['표고'] else '-'} m"
               + (f"   ⚠ {caution}" if caution else ""))
    c.font, c.fill, c.alignment = F_SUB, FILL_TITLE, AL_L
    ws.row_dimensions[2].height = 16
    r = 3
    ws.merge_cells(f"A{r}:E{r}")
    c = ws[f"A{r}"]
    c.value = f"소재지: {p['소재지']}"
    c.font, c.alignment = F_SM, AL_L
    # ── 결합키 — 이게 없으면 재방문·재시작·중심 이동을 구분할 수 없다 ──
    _lbl(ws, f"F{r}", "Visit ID")
    _fld(ws, f"G{r}")
    _lbl(ws, f"H{r}", "Location ID")
    _fld(ws, f"I{r}")
    ws[f"G{r}"].value = None
    ws[f"I{r}"].value = f"{p['_sid']}-P0a"
    ws[f"I{r}"].font = F_CALC
    ws[f"I{r}"].fill = FILL_CALC
    r += 1
    ws.merge_cells(f"A{r}:I{r}")
    c = ws[f"A{r}"]
    c.value = ("⚠ Visit ID 는 방문마다 새로 준다(V1·V2…). 중심점을 옮기면 **새 "
               "Location ID 로 새 카드**를 쓰고 이 카드 값을 덮지 않는다 — 옛 중심 "
               "자료가 새 중심을 입증하지 않기 때문이다.")
    c.font, c.fill, c.alignment = F_WARN, FILL_WARN, AL_TL
    ws.row_dimensions[r].height = 14
    r += 1

    # ── 0. 측정 전 확인 (§18③ 자기오염 방지) ────────────────
    #
    # ⚠️ 「지구물리측량 작업규정」 §18③ 은 측정자의 자기성 물품 제거를 명시한다.
    #    D·I 를 재지 않더라도 **총자력 측정에 직접 영향**을 준다 — 시계·펜·혁대
    #    하나가 네 방향 모두를 같은 방향으로 밀어 공간구배로 오인되게 만든다.
    r = _s(ws, r, "0. 측정 전 확인 — 자기오염 방지 (§18③)",
           "측정자 본인과 주변의 자기성 물품을 먼저 걷어낸다. 못 뺀 것이 있으면 "
           "무엇을 왜 못 뺐는지 적는다 — 나중에 그 방향 구배를 해석할 때 쓴다.")
    _lbl(ws, f"A{r}", "개인 금속물 제거")
    _fld(ws, f"B{r}")
    _dv(ws, f"B{r}", ["확인 — 전부 제거", "일부 미제거(사유 기재)"])
    _lbl(ws, f"C{r}", "차량 이격")
    _fld(ws, f"D{r}")
    _dv(ws, f"D{r}", ["확인 — 측정범위 밖", "근접(거리 기재)"])
    _lbl(ws, f"E{r}", "전자기기 차단")
    _fld(ws, f"F{r}")
    _dv(ws, f"F{r}", ["확인 — 전원 차단", "가동 중(사유 기재)"])
    _lbl(ws, f"G{r}", "미제거·예외 사유")
    _fld(ws, f"H{r}:I{r}")
    r += 2

    # ── 1. 도착·조건 ────────────────────────────────────────
    r = _s(ws, r, "1. 도착 · 관측 조건",
           "⚠ 절대측정을 하지 않으므로 «야간 의무»는 없다. 다만 시간변화가 안정된 "
           "구간을 골라야 한다 — Kp 와 우주기상 예보를 확인하고, 교란이 나면 그 "
           "구간은 빼거나 다시 잰다.")
    _lbl(ws, f"A{r}", "관측일자");   _fld(ws, f"B{r}", "yyyy-mm-dd")
    _lbl(ws, f"C{r}", "도착 시각");  _fld(ws, f"D{r}", "hh:mm")
    _lbl(ws, f"E{r}", "관측자");     _fld(ws, f"F{r}")
    _lbl(ws, f"G{r}", "기상");       _fld(ws, f"H{r}:I{r}")
    _dv(ws, f"H{r}", ["맑음", "흐림", "비", "눈", "바람 강함"])
    r += 1
    _lbl(ws, f"A{r}", "Kp 지수");    _fld(ws, f"B{r}", "0.0")
    _lbl(ws, f"C{r}", "예보 확인");  _fld(ws, f"D{r}")
    _dv(ws, f"D{r}", ["정온", "교란 예보", "미확인"])
    _lbl(ws, f"E{r}", "GSM-19T 기기번호"); _fld(ws, f"F{r}")
    _lbl(ws, f"G{r}", "Variometer");  _fld(ws, f"H{r}:I{r}")
    _dv(ws, f"H{r}", ["운용 — 연속기록", "미운용"])
    r += 2

    # ── 2. 기준점 좌표 ──────────────────────────────────────
    r = _s(ws, r, "2. 기준점 좌표 (GNSS 실측)",
           "십진도로 적는다. 사전좌표와 크게 다르면 사유를 소견에 남긴다.")
    _lbl(ws, f"A{r}", "위도");   _fld(ws, f"B{r}", "0.000000")
    _lbl(ws, f"C{r}", "경도");   _fld(ws, f"D{r}", "0.000000")
    _lbl(ws, f"E{r}", "취득방법"); _fld(ws, f"F{r}")
    _dv(ws, f"F{r}", ["RTK-GNSS", "네트워크RTK", "정적GNSS", "휴대GNSS"])
    _lbl(ws, f"G{r}", "사전값 대비(m)")
    _fld(ws, f"H{r}:I{r}", "0.0",
         calc=f'=IF(COUNT(B{r},D{r})<2,"",ROUND(SQRT((({p["경도"]}-D{r})*'
              f'111320*COS(RADIANS({p["위도"]})))^2+(({p["위도"]}-B{r})*110574)^2),1))')
    r += 2

    # ── 3. 수평 자기구배 ────────────────────────────────────
    r = _s(ws, r, "3. 수평 자기구배 — 방향마다 P0 재측정",
           f"한 방향을 다 잰 뒤 반드시 P0 로 돌아와 다시 잰다. 시각은 «시:분:초»까지 "
           f"적는다 — 사무실에서 이 시각으로 시간변화를 보정한다. "
           f"기준(참고): 반경 10 m 내 «4개 방위 측선의 관측점에서» "
           f"{SP.IAGA_RANGE_NT:.0f} nT · 구배 {SP.EURO_GRAD_NT_PER_M:.0f} nT/m "
           f"— 국제 권고이고 국내 기준은 미확정이다. 측선 사이와 대각선은 재지 "
           f"않으므로 「반경 10 m 전체」를 말할 수 없다. 판정은 사무실에서 한다.")
    _lbl(ws, f"A{r}", "Set ID 접두")
    _fld(ws, f"B{r}:C{r}", calc=f'="{p["_sid"]}-H"&"(방향)"')
    _lbl(ws, f"D{r}", "차수(Attempt)")
    _fld(ws, f"E{r}")
    _dv(ws, f"E{r}", ["1", "2", "3"])
    _lbl(ws, f"F{r}", "실측거리 사용")
    _fld(ws, f"G{r}")
    _dv(ws, f"G{r}", ["명목거리 그대로", "실측(비고에 기재)"])
    _lbl(ws, f"H{r}:I{r}", "중단 시 → 차수 올려 새 Set")
    r += 1
    hdr_r = r
    _lbl(ws, f"A{r}", "거리(m)")
    for k, d in enumerate(SP.H_DIRECTIONS):
        col1, col2 = COLS[1 + k * 2], COLS[2 + k * 2]
        _lbl(ws, f"{col1}{r}", f"{d} 시각")
        _lbl(ws, f"{col2}{r}", f"{d} F(nT)")
    ws.row_dimensions[r].height = 18
    r += 1
    first_data = r
    seq = ["P0(전)"] + [f"{x} m" for x in SP.H_OFFSETS_M] + ["P0(후)"]
    rows_of = {}
    for lab in seq:
        isP0 = lab.startswith("P0")
        c = ws[f"A{r}"]
        c.value, c.font, c.border, c.alignment = lab, F_LBL, BOX, AL_C
        c.fill = FILL_P0 if isP0 else FILL_LBL
        for k in range(len(SP.H_DIRECTIONS)):
            _fld(ws, f"{COLS[1+k*2]}{r}", "hh:mm:ss")
            _fld(ws, f"{COLS[2+k*2]}{r}", "0.0")
            if isP0:
                ws[f"{COLS[1+k*2]}{r}"].fill = FILL_P0
                ws[f"{COLS[2+k*2]}{r}"].fill = FILL_P0
        rows_of[lab] = r
        r += 1
    last_data = r - 1
    # 현장 즉시 확인용 두 지표 — 정밀 보정은 사무실에서
    p0a, p0b = rows_of["P0(전)"], rows_of["P0(후)"]
    m_lo = rows_of[f"{SP.H_OFFSETS_M[0]} m"]
    m_hi = rows_of[f"{SP.H_OFFSETS_M[-1]} m"]
    for lab, fml in (
        ("P0 전후차(nT)",
         lambda c: f'=IF(COUNT({c}{p0a},{c}{p0b})<2,"",{c}{p0b}-{c}{p0a})'),
        ("관측값 max−min(nT)",
         lambda c: f'=IF(COUNT({c}{m_lo}:{c}{m_hi})=0,"",'
                   f'MAX({c}{m_lo}:{c}{m_hi})-MIN({c}{m_lo}:{c}{m_hi}))'),
    ):
        _lbl(ws, f"A{r}", lab)
        for k in range(len(SP.H_DIRECTIONS)):
            col = COLS[2 + k * 2]
            _lbl(ws, f"{COLS[1+k*2]}{r}", "")
            _fld(ws, f"{col}{r}", "0.0", calc=fml(col))
        r += 1
    ws.merge_cells(f"A{r}:I{r}")
    c = ws[f"A{r}"]
    c.value = ("⚠ P0 전후차가 크면(수십 nT) 그 구간에 자기장이 크게 흔들린 것이다 — "
               "다시 재거나 소견에 적는다. 위 두 값은 «보정 전» 참고치이며 판정값이 "
               "아니다.  ※ 판독은 위치당 1회(기기 내부평균) 기준이다. 값이 의심스러우면 "
               "차수를 올려 그 방향을 통째로 다시 잰다.")
    c.font, c.fill, c.alignment = F_WARN, FILL_WARN, AL_TL
    ws.row_dimensions[r].height = 26
    r += 1
    # ── 현장 완전성 점검 (Codex M3) — 자동지표만으로는 못 정한다 ──
    _lbl(ws, f"A{r}", "완전성 점검")
    for lab, col in (("필수행 다 채움", "B"), ("시각 시:분:초", "C"),
                     ("P0 전·후 있음", "D"), ("측정시각 P0 사이", "E"),
                     ("Variometer 연결", "F")):
        _lbl(ws, f"{col}{r}", lab)
    _lbl(ws, f"G{r}:I{r}", "하나라도 아니오 → 철수 전 보완")
    r += 1
    _lbl(ws, f"A{r}", "예 / 아니오")
    for col in "BCDEF":
        _fld(ws, f"{col}{r}")
        _dv(ws, f"{col}{r}", ["예", "아니오"])
    _fld(ws, f"G{r}:I{r}")
    r += 2

    # ── 4. 수직 자기구배 ────────────────────────────────────
    r = _s(ws, r, "4. 수직 자기구배 — 중심점(P0)에서 높이별",
           "기준높이에서 시작·종료 두 번 잰다. 실제 센서 «중심» 높이를 적는다 — "
           "목표 높이만으로는 재현되지 않는다.")
    _lbl(ws, f"A{r}", "Set ID")
    _fld(ws, f"B{r}", calc=f'="{p["_sid"]}-V"')
    _lbl(ws, f"C{r}", "차수")
    _fld(ws, f"D{r}")
    _dv(ws, f"D{r}", ["1", "2", "3"])
    _lbl(ws, f"E{r}", "측정 기준면")
    _fld(ws, f"F{r}")
    _dv(ws, f"F{r}", ["지면", "표석 상면"])
    _lbl(ws, f"G{r}", "기준높이(cm)")
    _fld(ws, f"H{r}:I{r}", "0")
    r += 1
    _lbl(ws, f"A{r}", "명목 높이(cm)")
    _lbl(ws, f"B{r}", "실제 높이(cm)")
    _lbl(ws, f"C{r}", "시각")
    _lbl(ws, f"D{r}", "F(nT)")
    _lbl(ws, f"E{r}:I{r}", "비고")
    r += 1
    vseq = ["기준(전)"] + [f"{h} cm" for h in SP.V_HEIGHTS_CM] + ["기준(후)"]
    v_first = r
    for lab in vseq:
        isR = lab.startswith("기준")
        c = ws[f"A{r}"]
        c.value, c.font, c.border, c.alignment = lab, F_LBL, BOX, AL_C
        c.fill = FILL_P0 if isR else FILL_LBL
        _fld(ws, f"B{r}", "0.0")
        _fld(ws, f"C{r}", "hh:mm:ss")
        _fld(ws, f"D{r}", "0.0")
        _fld(ws, f"E{r}:I{r}")
        if isR:
            for cc in "BCD":
                ws[f"{cc}{r}"].fill = FILL_P0
        r += 1
    _lbl(ws, f"A{r}", "F max−min(nT)")
    _fld(ws, f"B{r}", calc="")
    _fld(ws, f"C{r}", calc="")
    _fld(ws, f"D{r}", "0.0",
         calc=f'=IF(COUNT(D{v_first+1}:D{r-2})=0,"",'
              f'MAX(D{v_first+1}:D{r-2})-MIN(D{v_first+1}:D{r-2}))')
    _fld(ws, f"E{r}:I{r}", calc='="측정 높이 구간의 총자기장 변화폭 (참고)"')
    r += 2

    # ── 5. 방위표지 ─────────────────────────────────────────
    r = _s(ws, r, "5. 방위표지 시준 — 기준점 → 표지 방향 (진북 · 북 0° · 시계방향)",
           "원시각은 곤(gon, 1회전 400 · 대척 200)으로 적는다. 폐합차가 0.02 gon"
           "(약 65″)을 넘으면 다시 시준한다. ⚠ 폐합차가 작아도 «참방위각이 틀릴 수» "
           "있다 — 그래서 표지 좌표·취득방법·기존 확정값을 함께 남긴다. 재방문 잔여 "
           "RMS 33.7′ 의 원인이 바로 방문마다 달라진 참방위각이었다.")
    _lbl(ws, f"A{r}", "표지");         _lbl(ws, f"B{r}", "Mark ID")
    _lbl(ws, f"C{r}:D{r}", "시준 지점(표지의 어디)")
    _lbl(ws, f"E{r}", "표지 위도");    _lbl(ws, f"F{r}", "표지 경도")
    _lbl(ws, f"G{r}", "좌표 취득방법"); _lbl(ws, f"H{r}", "정확도(m)")
    _lbl(ws, f"I{r}", "기존 확정값(°)")
    r += 1
    for k in (1, 2):
        c = ws[f"A{r}"]
        c.value, c.font, c.fill, c.border, c.alignment = (
            f"방위표지{k}", F_LBL, FILL_LBL, BOX, AL_C)
        _fld(ws, f"B{r}", calc=f'="{p["_sid"]}-M{k}"')
        _fld(ws, f"C{r}:D{r}")
        for cc, fm in (("E", "0.000000"), ("F", "0.000000"), ("H", "0.00"),
                       ("I", "0.0000")):
            _fld(ws, f"{cc}{r}", fm)
        _fld(ws, f"G{r}")
        _dv(ws, f"G{r}", ["RTK-GNSS", "네트워크RTK", "정적GNSS", "휴대GNSS",
                          "기존 성과 인용"])
        r += 1
    _lbl(ws, f"A{r}", "표지");        _lbl(ws, f"B{r}", "정 시준(gon)")
    _lbl(ws, f"C{r}", "반 시준(gon)"); _lbl(ws, f"D{r}", "폐합차(gon)")
    _lbl(ws, f"E{r}", "참방위각(gon)"); _lbl(ws, f"F{r}", "결정방법")
    _lbl(ws, f"G{r}:I{r}", "표지 변경 여부 · 사유")
    r += 1
    for k in (1, 2):
        c = ws[f"A{r}"]
        c.value, c.font, c.fill, c.border, c.alignment = (
            f"방위표지{k}", F_LBL, FILL_LBL, BOX, AL_C)
        _fld(ws, f"B{r}", "0.0000")
        _fld(ws, f"C{r}", "0.0000")
        _fld(ws, f"D{r}", "0.0000",
             calc=f'=IF(COUNT(B{r}:C{r})<2,"",ABS(ABS(C{r}-B{r})-200))')
        _fld(ws, f"E{r}", "0.0000")
        _fld(ws, f"F{r}")
        _dv(ws, f"F{r}", ["천문관측", "자이로", "RTK 장기선", "좌표계산",
                          "기존 성과 인용"])
        _fld(ws, f"G{r}:I{r}")
        r += 1
    r += 1

    # ── 6. 현장 사진 ────────────────────────────────────────
    #
    # ⚠️ 사진은 «무엇을 찍었는지»가 표에 남아야 쓸모가 있다. 파일명을 적는
    #    칸을 상자마다 붙인 것은 그래서다 — 엑셀에 붙인 그림은 나중에 파일과
    #    대조하기 어렵고, 용량 때문에 빠지는 일도 있다.
    r = _s(ws, r, "6. 현장 사진",
           "상자 안에 사진을 붙이고 아래 칸에 «원본 파일명»과 촬영시각을 적는다. "
           f"명명규칙은 {{{p['_sid']}}}_{{VisitID}}_{{코드}}_{{연번}} 이다"
           "(코드: P0·MEAS·M1·M2·E·W·S·N·INT). ⚠ 파일명만으로는 카메라 간 중복·"
           "이름 변경·복사 누락을 잡지 못하므로 촬영시각을 함께 남긴다. "
           "방위별 전경은 «중심점에 서서» 그 방향을 보고, **측선과 10 m 종점이 "
           "보이게** 찍는다 — 구배가 큰 방향의 원인을 사진으로 되짚는다. "
           "자기교란 요소는 P0 와의 방향·거리가 보이는 관계사진과 근접사진을 "
           "함께 담고, 여러 개면 소견에 목록을 적는다.")
    SPANS = [("A", "C"), ("D", "F"), ("G", "I")]
    SHOTS = [
        [("중심점(P0) 전경", "표석·마커가 보이게"),
         ("측정 모습", "센서 높이를 알 수 있게"),
         ("방위표지1", "시준 지점이 보이게")],
        [("동(E) 방향 전경", "중심점에서 동쪽을 보고"),
         ("서(W) 방향 전경", "중심점에서 서쪽을 보고"),
         ("남(S) 방향 전경", "중심점에서 남쪽을 보고")],
        [("북(N) 방향 전경", "중심점에서 북쪽을 보고"),
         ("방위표지2", "시준 지점이 보이게"),
         ("자기교란 요소 · 기타", "없으면 「해당없음」")],
    ]
    for band in SHOTS:
        for (a, b), (t, hint) in zip(SPANS, band):
            _lbl(ws, f"{a}{r}:{b}{r}", t)
        r += 1
        for (a, b), (t, hint) in zip(SPANS, band):
            _photo(ws, r, (a, b), f"[ {hint} ]")
        r += 8
        for (a, b), (t, hint) in zip(SPANS, band):
            _lbl(ws, f"{a}{r}", "파일명")
            _fld(ws, f"{chr(ord(a)+1)}{r}:{b}{r}")
        r += 1
        for (a, b), (t, hint) in zip(SPANS, band):
            _lbl(ws, f"{a}{r}", "촬영시각")
            _fld(ws, f"{chr(ord(a)+1)}{r}:{b}{r}", "hh:mm:ss")
        r += 1
    r += 1

    # ── 7. 중심점 최종 · 소견 ───────────────────────────────
    r = _s(ws, r, "7. 중심점 최종 결정 · 현장 소견",
           "구배가 큰 방향이 있으면 후보지 안에서 더 조용한 자리로 중심점을 옮길 수 "
           "있다. **옮겼으면 옮긴 자리에서 3·4 를 다시 재야 한다.**")
    _lbl(ws, f"A{r}", "중심점 조정"); _fld(ws, f"B{r}")
    _dv(ws, f"B{r}", ["조정 없음", "조정함"])
    _lbl(ws, f"C{r}", "최종 위도");  _fld(ws, f"D{r}", "0.000000")
    _lbl(ws, f"E{r}", "최종 경도");  _fld(ws, f"F{r}", "0.000000")
    _lbl(ws, f"G{r}", "조정 후 재측정"); _fld(ws, f"H{r}:I{r}")
    _dv(ws, f"H{r}", ["재측정 완료", "미실시 — 잠정", "해당없음"])
    r += 1
    _lbl(ws, f"A{r}", "표지 설치");  _fld(ws, f"B{r}")
    _dv(ws, f"B{r}", ["설치 완료", "미설치"])
    _lbl(ws, f"C{r}", "철수 시각");  _fld(ws, f"D{r}", "hh:mm")
    _lbl(ws, f"E{r}", "사진 매수");  _fld(ws, f"F{r}", "0")
    _lbl(ws, f"G{r}", "현장 판정(소견)"); _fld(ws, f"H{r}:I{r}")
    _dv(ws, f"H{r}", ["양호", "보통", "불량 — 대체지 필요", "재측정 필요"])
    r += 1
    _lbl(ws, f"A{r}", "현장 소견")
    ws.merge_cells(f"B{r}:I{r+3}")
    c = ws[f"B{r}"]
    c.font, c.fill, c.alignment, c.border = F_VAL, FILL_FIELD, AL_TL, BOX
    for rr in range(r, r + 4):
        for cc in COLS:
            ws[f"{cc}{rr}"].border = BOX
        ws.row_dimensions[rr].height = 16
    r += 4
    ws.print_area = f"A1:I{r}"
    return ws


# ══════════════════════════════════════════════════════════════
def build_master(wb, pts):
    ws = wb.create_sheet("총괄", 0)
    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = f"지자기 시험 탐사 현장야장 — 대상 {len(pts)}점"
    c.font, c.fill, c.alignment = F_TITLE, FILL_TITLE, AL_L
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:K2")
    c = ws["A2"]
    c.value = ("점마다 시트가 하나씩 있다. 현장에서는 그 시트만 채우면 된다. "
               "시간변화 보정·구배 산출·판정은 사무실에서 「시험탐사 표준야장」으로 한다.")
    c.font, c.alignment = F_SM, AL_L
    hdr = ["연번", "Site ID", "구분", "지점명", "도엽명", "위도", "경도",
           "유의사항", "관측일자", "관측자", "진행 상태"]
    for j, h in enumerate(hdr, 1):
        cc = ws.cell(row=4, column=j, value=h)
        cc.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        cc.fill, cc.alignment, cc.border = FILL_TITLE, AL_C, BOX
    for j, w in enumerate([5, 10, 7, 18, 11, 11, 11, 20, 12, 11, 15], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    for i, p in enumerate(pts, 1):
        r = 4 + i
        vals = [i, p["_sid"], p["구분"], p["지점명"], p["도엽명"], p["위도"],
                p["경도"], SP.SITE_CAUTION.get(p["지점명"], ""), None, None, None]
        for j, v in enumerate(vals, 1):
            cc = ws.cell(row=r, column=j, value=v)
            cc.font, cc.border = F_VAL, BOX
            cc.alignment = AL_L if j == 8 else AL_C
            cc.fill = FILL_FIELD if j >= 9 else PatternFill("solid", fgColor="FFFFFF")
            if j in (6, 7):
                cc.number_format = "0.000000"
            if j == 9:
                cc.number_format = "yyyy-mm-dd"
        if SP.SITE_CAUTION.get(p["지점명"]):
            ws.cell(row=r, column=8).font = F_WARN
    last = 4 + len(pts)
    _dv(ws, f"K5:K{last}", ["미착수", "진행 중", "완료", "재방문 필요"])
    return ws


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    pts = TP.load_points()
    for i, p in enumerate(pts, 1):
        p["_sid"] = f"S{i:02d}"
    wb = Workbook()
    wb.remove(wb.active)
    for i, p in enumerate(pts, 1):
        build_card(wb, p, i)
    build_master(wb, pts)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{dt.datetime.now():%Y%m%d_%H%M%S}_시험탐사_현장야장.xlsx"
    wb.save(path)
    print(f"카드 {len(pts)}장 + 총괄 1 = 시트 {len(wb.sheetnames)}장")
    print(f"[저장] {path}  ({path.stat().st_size/1e6:.2f} MB)")
    return path


if __name__ == "__main__":
    main()
