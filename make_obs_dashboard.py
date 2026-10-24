# -*- coding: utf-8 -*-
"""
기존 측정점 자료 점검 대시보드 — HTML · 엑셀 (로컬 전용)
=========================================================

```
python build_obs_archive.py     # ① 다섯 출처 통합
python lmm_intake_decide.py     # ② 판정
python make_obs_dashboard.py    # ③ 대시보드 → local/
```

⚠️ **`local/` 은 `.gitignore` 로 막혀 있다.** 내부 점검용이라 배포하지 않는다.
   산출물에 원본 자료의 **절대 경로**(D: 지리원 자료)가 들어가므로 공개하면
   안 되는 것도 이유다.

설계 근거 — 어떤 대시보드가 이 일에 맞는가
-------------------------------------------
점검용 대시보드의 고전적 설계 원칙 몇 가지를 그대로 따랐다.

  **① 개요 먼저, 자세한 것은 요청할 때**(Shneiderman 의 정보탐색 원칙)
     지도와 측점 목록이 한눈에 전부를 보이고, 측점을 고르면 그때 시계열·
     이력·플래그가 열린다. 처음부터 587행을 펼치지 않는다.

  **② 한 화면 · 장식 없음**(Few 의 대시보드 설계)
     머리의 숫자 띠 → 지도 → 상세. 스크롤 없이 전체 상태가 읽히게 한다.

  **③ 연결된 선택**(brushing & linking)
     지도·목록·시계열·플래그가 모두 「지금 고른 측점」 하나를 공유한다.
     어디서 고르든 나머지가 따라온다.

  **④ 성분 탭과 공통 시간축**(INTERMAGNET·USGS 관측소 뷰어의 어법)
     편각·복각·총자력을 **세로로 쌓아** 같은 연도축에 둔다. 한 해에 세
     성분이 함께 튀었는지가 바로 보인다.

  **⑤ 품질 플래그는 심각도 순 + 원본으로 바로 가기**(자료품질 대시보드)
     플래그마다 「무엇이·얼마나·왜·어느 파일을 열면 되는가」 네 가지를 적는다.
     ④가 없으면 점검을 시작할 수 없다.

이 자료에 맞춘 선택 둘
----------------------
  **잔차를 기본으로 그린다.** 편각 실측값은 영년변화 때문에 저절로 기울어
  있어 그래프가 「변하고 있다」만 보여 준다. **IGRF 잔차**는 영년변화가
  빠진 양이라 같은 표석이면 평평해야 하고, 흔들리면 그 방문에 문제가 있다는
  뜻이다. 실측값은 단추 하나로 바꿔 본다.

  **출처를 색으로 가른다.** 같은 해를 두 문서가 다르게 적은 일이 잦아,
  점검에서는 「누가 무엇을 적었나」가 값 자체만큼 중요하다.

⚠️ **자동으로 버리지 않는다.** 플래그는 사람이 확인할 목록이지 배제 목록이
   아니다(이 저장소의 「현장 탈락 금지」 규약과 같다).
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).parent
OUT = ROOT / "local"
BOUND = ROOT / "data" / "korea_boundary.geojson"

import obs_dashboard_data as ODD                     # noqa: E402

SEV_COLOR = {"확인필요": "#B2182B", "주의": "#E08214", "참고": "#4575B4"}
SRC_COLOR = {
    "점조서": "#6A51A3",
    "관측현황 10~19": "#2171B5",
    "성과표 22~25": "#238B45",
    "절대관측 원자료 12~16": "#D94801",
    "야장 19~25": "#CE1256",
}


def jnum(v, nd=4):
    """JSON 에 넣을 수 있는 수로. NaN·inf 는 None 이다."""
    try:
        x = float(v)
        return None if not np.isfinite(x) else round(x, nd)
    except Exception:
        return None


def coastline():
    """해안선을 [[ [lon,lat], … ], …] 로. 외부 타일을 쓰지 않는다."""
    if not BOUND.exists():
        return []
    j = json.loads(BOUND.read_text(encoding="utf-8"))
    rings = []
    for ft in j.get("features", []):
        g = ft.get("geometry", {})
        t, c = g.get("type"), g.get("coordinates")
        polys = [c] if t == "Polygon" else (c if t == "MultiPolygon" else [])
        for poly in polys:
            if poly:
                # 바깥 고리만 — 섬 안의 호수까지 그릴 필요가 없다
                rings.append([[round(p[0], 4), round(p[1], 4)] for p in poly[0]])
    return rings


def payload():
    ST, V, S, F = ODD.build()

    stations = []
    for _, r in ST.iterrows():
        stations.append({k: (jnum(r[k], 6) if k in ("위도", "경도")
                             else jnum(r[k]) if k in
                             ("rD_중앙", "rD_폭", "rF_중앙", "rF_폭", "같은날산포")
                             else (None if pd.isna(r[k]) else
                                   (int(r[k]) if k in
                                    ("방문", "최초", "최종", "출처수", "투입D",
                                     "플래그", "확인필요") else str(r[k]))))
                         for k in ST.columns})

    vis = []
    for _, r in V.sort_values(["측점", "연도", "출처"]).iterrows():
        vis.append(dict(
            측점=str(r["측점"]), 연도=(int(r["연도"]) if pd.notna(r["연도"]) else None),
            일자=str(r["일자"]), 출처=str(r["출처명"]), 세션=int(r["세션"]),
            D=jnum(r["D"], 6), I=jnum(r["I"], 6), F=jnum(r["F"], 3),
            rD=jnum(r["rD분"], 2), rI=jnum(r["rI분"], 2), rF=jnum(r["rF_nT"], 1),
            D산포=jnum(r.get("D_산포"), 2), I산포=jnum(r.get("I_산포"), 2),
            F산포=jnum(r.get("F_산포"), 1),
            D판정=str(r.get("D판정", "")), I판정=str(r.get("I판정", "")),
            F판정=str(r.get("F판정", "")),
            좌표신뢰=str(r.get("좌표신뢰", "") or ""),
            일자정밀도=str(r.get("일자정밀도", "") or ""),
            원본파일=str(r.get("원본파일", "") or ""),
            원본시트=str(r.get("원본시트", "") or ""),
            참방위각=jnum(r.get("참방위각"), 4)))

    ses = []
    for _, r in S.sort_values(["측점", "연도", "출처", "회차"]).iterrows():
        ses.append(dict(
            측점=str(r["측점"]), 연도=(int(r["연도"]) if pd.notna(r["연도"]) else None),
            일자=str(r["일자"]), 출처=str(r["출처명"]),
            회차=(int(r["회차"]) if pd.notna(r["회차"]) else None),
            D=jnum(r["D"], 6), I=jnum(r["I"], 6), F=jnum(r["F"], 3),
            rD=jnum(r["rD분"], 2), rI=jnum(r["rI분"], 2), rF=jnum(r["rF_nT"], 1),
            판정=f"{r['D_판정']}/{r['I_판정']}/{r['F_판정']}",
            원본파일=str(r.get("원본파일", "") or ""),
            원본시트=str(r.get("원본시트", "") or "")))

    fl = []
    for _, r in F.iterrows():
        fl.append({k: (None if pd.isna(r[k]) else
                       (int(r[k]) if k == "연도" else str(r[k])))
                   for k in F.columns})

    # 원본 자료 — 파일마다 한 줄
    src = (S[S["원본파일"].astype(str).str.len() > 0]
           .groupby(["출처명", "원본파일"])
           .agg(측점=("측점", "nunique"), 행=("측점", "size"),
                최초=("연도", "min"), 최종=("연도", "max"),
                측점목록=("측점", lambda x: ", ".join(sorted(set(x)))))
           .reset_index())
    srcs = []
    for _, r in src.sort_values(["출처명", "원본파일"]).iterrows():
        p = Path(str(r["원본파일"]))
        srcs.append(dict(출처=str(r["출처명"]), 파일=p.name, 폴더=str(p.parent),
                         전체=str(p), 있음=bool(p.exists()),
                         측점=int(r["측점"]), 행=int(r["행"]),
                         최초=(int(r["최초"]) if pd.notna(r["최초"]) else None),
                         최종=(int(r["최종"]) if pd.notna(r["최종"]) else None),
                         측점목록=str(r["측점목록"])))

    return dict(생성=f"{dt.datetime.now():%Y-%m-%d %H:%M}",
                측점=stations, 방문=vis, 세션=ses, 플래그=fl,
                원본=srcs, 해안선=coastline(),
                색_심각도=SEV_COLOR, 색_출처=SRC_COLOR), (ST, V, S, F)


# ────────────────────────────────────────────────────────────────────
#  HTML
# ────────────────────────────────────────────────────────────────────
CSS = """
*{box-sizing:border-box}
body{margin:0;font:13px/1.55 "맑은 고딕","Malgun Gothic",system-ui,sans-serif;
 color:#1A1A1A;background:#F4F5F7}
header{background:#16324F;color:#fff;padding:10px 16px;display:flex;
 align-items:baseline;gap:14px;flex-wrap:wrap}
header h1{margin:0;font-size:16px;font-weight:700;letter-spacing:-.2px}
header .sub{font-size:11px;opacity:.78}
.kpis{display:flex;gap:8px;margin-left:auto;flex-wrap:wrap}
.kpi{background:rgba(255,255,255,.12);border-radius:4px;padding:4px 10px;
 text-align:center;cursor:default;min-width:72px}
.kpi b{display:block;font-size:17px;line-height:1.15}
.kpi span{font-size:10px;opacity:.82}
.kpi.act{cursor:pointer}.kpi.act:hover{background:rgba(255,255,255,.26)}
.kpi.bad b{color:#FFB3B3}
main{display:grid;grid-template-columns:330px 1fr;gap:10px;padding:10px;
 height:calc(100vh - 52px)}
.pane{background:#fff;border:1px solid #D7DBE0;border-radius:5px;
 display:flex;flex-direction:column;min-height:0;min-width:0}
.pane>h2{margin:0;padding:7px 10px;font-size:12px;font-weight:700;
 border-bottom:1px solid #E4E7EB;background:#FAFBFC;color:#16324F;
 display:flex;align-items:center;gap:8px}
.pane>h2 .r{margin-left:auto;font-weight:400;color:#6B7280;font-size:11px}
#left{display:grid;grid-template-rows:auto 1fr;gap:10px;min-height:0}
#mapwrap{position:relative}
#map{display:block;width:100%;height:auto}
#listbox{overflow:auto}
table{border-collapse:collapse;width:100%;font-size:12px}
th,td{padding:3px 7px;border-bottom:1px solid #EDEFF2;text-align:left;
 white-space:nowrap}
th{position:sticky;top:0;background:#F7F8FA;font-weight:700;font-size:11px;
 color:#4B5563;cursor:pointer;user-select:none;z-index:2}
th:hover{background:#EEF1F5}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
tbody tr{cursor:pointer}
tbody tr:hover{background:#F0F6FF}
tbody tr.sel{background:#DCEBFF;box-shadow:inset 3px 0 0 #16324F}
.tabs{display:flex;gap:2px;padding:6px 8px 0;border-bottom:1px solid #E4E7EB;
 background:#FAFBFC;flex-wrap:wrap}
.tab{padding:5px 12px;font-size:12px;border:1px solid transparent;
 border-bottom:none;border-radius:4px 4px 0 0;cursor:pointer;color:#4B5563}
.tab:hover{background:#EEF1F5}
.tab.on{background:#fff;border-color:#D7DBE0;color:#16324F;font-weight:700;
 margin-bottom:-1px}
#body{overflow:auto;padding:10px;flex:1;min-height:0}
.chart{width:100%;height:auto;display:block}
.sev{display:inline-block;padding:1px 6px;border-radius:9px;color:#fff;
 font-size:10px;font-weight:700}
.dot{display:inline-block;width:8px;height:8px;border-radius:50%;
 margin-right:5px;vertical-align:-1px}
.muted{color:#6B7280}
.warn{color:#B2182B;font-weight:700}
.tools{display:flex;gap:6px;align-items:center;padding:5px 8px;
 border-bottom:1px solid #EDEFF2;flex-wrap:wrap;background:#fff}
input[type=search],select{font:12px inherit;padding:3px 6px;
 border:1px solid #CBD2D9;border-radius:3px;background:#fff}
input[type=search]{flex:1;min-width:90px}
button{font:12px inherit;padding:3px 9px;border:1px solid #CBD2D9;
 border-radius:3px;background:#fff;cursor:pointer}
button:hover{background:#F0F6FF}
button.on{background:#16324F;color:#fff;border-color:#16324F}
.legend{display:flex;gap:10px;flex-wrap:wrap;font-size:11px;padding:4px 8px;
 color:#4B5563;border-bottom:1px solid #EDEFF2}
.path{font-family:Consolas,monospace;font-size:11px;color:#374151;
 word-break:break-all;white-space:normal}
.note{background:#FFF8E6;border-left:3px solid #E08214;padding:6px 9px;
 font-size:11px;margin:0 0 9px;border-radius:0 3px 3px 0}
.note.bad{background:#FDECEC;border-left-color:#B2182B}
h3{margin:14px 0 5px;font-size:12px;color:#16324F}
h3:first-child{margin-top:0}
#tip{position:fixed;display:none;background:#1A1A1A;color:#fff;padding:5px 8px;
 border-radius:3px;font-size:11px;pointer-events:none;z-index:99;max-width:300px}
@media(max-width:900px){main{grid-template-columns:1fr;height:auto}
 #left{grid-template-rows:auto auto}#listbox{max-height:300px}}
"""


def build_html(P):
    data = json.dumps(P, ensure_ascii=False, separators=(",", ":"))
    js = (ROOT / "_obs_dashboard.js")
    script = js.read_text(encoding="utf-8") if js.exists() else JS
    n_sev = {}
    for f in P["플래그"]:
        n_sev[f["심각도"]] = n_sev.get(f["심각도"], 0) + 1
    return f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>기존 측정점 자료 점검</title>
<style>{CSS}</style></head><body>
<header>
  <h1>기존 측정점 자료 점검</h1>
  <div class="sub">1등 지자기점 관측망 · {P['생성']} 생성 · 내부 점검용</div>
  <div class="kpis">
    <div class="kpi"><b>{len(P['측점'])}</b><span>측점</span></div>
    <div class="kpi"><b>{len(P['방문'])}</b><span>방문</span></div>
    <div class="kpi"><b>{len(P['세션'])}</b><span>세션</span></div>
    <div class="kpi"><b>{len(P['원본'])}</b><span>원본 파일</span></div>
    <div class="kpi act bad" onclick="allFlags('확인필요')">
      <b>{n_sev.get('확인필요', 0)}</b><span>확인 필요</span></div>
    <div class="kpi act" onclick="allFlags('주의')">
      <b>{n_sev.get('주의', 0)}</b><span>주의</span></div>
    <div class="kpi act" onclick="allFlags('')">
      <b>{len(P['플래그'])}</b><span>전체 플래그</span></div>
  </div>
</header>
<main>
  <div id="left">
    <div class="pane"><h2>관측망 지도<span class="r" id="maphint">측점을 누르세요</span></h2>
      <div class="legend" id="maplegend"></div>
      <div id="mapwrap"><canvas id="map" width="640" height="840"></canvas></div>
    </div>
    <div class="pane"><h2>측점 <span class="r" id="listcount"></span></h2>
      <div class="tools">
        <input type="search" id="q" placeholder="측점 이름 검색">
        <select id="filt">
          <option value="">전체</option>
          <option value="flag">확인 필요가 있는 측점</option>
          <option value="lmm">현재 모델이 쓰는 측점</option>
          <option value="nolmm">모델에 없는 측점</option>
        </select>
      </div>
      <div id="listbox"><table id="list"></table></div>
    </div>
  </div>
  <div class="pane">
    <h2 id="title">측점을 고르세요<span class="r" id="subtitle"></span></h2>
    <div class="tabs" id="tabs"></div>
    <div id="body"></div>
  </div>
</main>
<div id="tip"></div>
<script id="DATA" type="application/json">{data}</script>
<script>{script}</script>
</body></html>"""


JS = ""   # 아래 _obs_dashboard.js 로 분리한다


def build_excel(ST, V, S, F, P, path):
    """엑셀 판 — 같은 자료를 피벗·필터로 볼 수 있게 평탄하게 낸다."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    HEAD = PatternFill("solid", fgColor="16324F")
    HF = Font(color="FFFFFF", bold=True, size=10, name="맑은 고딕")
    BF = Font(size=10, name="맑은 고딕")
    SEVFILL = {"확인필요": PatternFill("solid", fgColor="FADBD8"),
               "주의": PatternFill("solid", fgColor="FDEBD0"),
               "참고": PatternFill("solid", fgColor="EAF2F8")}

    def sheet(name, df, widths=None, sevcol=None, freeze="A2"):
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(1, c)
            cell.fill, cell.font = HEAD, HF
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
        for _, r in df.iterrows():
            ws.append([None if (isinstance(v, float) and not np.isfinite(v))
                       else v for v in r.tolist()])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = BF
        if sevcol and sevcol in df.columns:
            ci = list(df.columns).index(sevcol) + 1
            for i in range(2, len(df) + 2):
                f = SEVFILL.get(str(ws.cell(i, ci).value))
                if f:
                    for c in range(1, len(df.columns) + 1):
                        ws.cell(i, c).fill = f
        ws.freeze_panes = freeze
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(df.columns))}"
                              f"{len(df) + 1}")
        for i, col in enumerate(df.columns, 1):
            w = (widths or {}).get(col)
            if w is None:
                vals = [str(col)] + [str(x)[:60] for x in df[col].head(220)]
                w = min(46, max(7, int(max(len(x) for x in vals) * 1.15) + 2))
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    # ── ① 개요 ────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "① 개요"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 76
    rows = [
        ("기존 측정점 자료 점검 대시보드", "", ""),
        (f"{P['생성']} 생성 · 내부 점검용", "", ""),
        ("", "", ""),
        ("측점", len(ST), "1등 지자기점 관측망"),
        ("방문", len(V), "(측점, 일자, 출처) 조합. 같은 해를 두 출처가 적으면 둘 다 센다"),
        ("세션", len(S), "회차별 원시 판독값까지"),
        ("원본 파일", len(P["원본"]), "⑥ 시트에 전체 경로"),
        ("기간", f"{int(ST['최초'].min())}~{int(ST['최종'].max())}", ""),
        ("", "", ""),
        ("확인 필요", int((F["심각도"] == "확인필요").sum()), "⑤ 시트 — 먼저 볼 것"),
        ("주의", int((F["심각도"] == "주의").sum()), ""),
        ("참고", int((F["심각도"] == "참고").sum()), ""),
        ("", "", ""),
        ("■ 시트 안내", "", ""),
        ("② 측점", "", "측점마다 한 줄 — 기간·방문수·잔차 변동·플래그 수"),
        ("③ 방문 시계열", "", "연도별 D·I·F 와 IGRF 잔차. 피벗·그래프의 바탕"),
        ("④ 세션 원자료", "", "회차별 판독값. 가장 원본에 가깝다"),
        ("⑤ 확인 필요", "", "심각도 순. 무엇이·얼마나·왜·어느 파일인지"),
        ("⑥ 원본 자료 위치", "", "파일마다 전체 경로와 담긴 측점"),
        ("⑦ 잔차 그래프", "", "측점별 편각 잔차 — 같은 표석이면 평평해야 한다"),
        ("", "", ""),
        ("■ 읽는 법", "", ""),
        ("IGRF 잔차", "", "영년변화가 빠진 양이다. 같은 표석을 다시 재면 같아야 하고, "
                        "흔들리면 그 방문에 문제가 있다는 뜻이다"),
        ("실측값", "", "편각은 영년변화로 저절로 기울어 있어 그래프가 "
                     "「변하고 있다」만 보여 준다"),
        ("플래그", "", "사람이 확인할 목록이지 «배제 목록»이 아니다. "
                     "자동으로 버리지 않는다"),
    ]
    for a, b, c in rows:
        ws.append([a, b, c])
    for i in range(1, ws.max_row + 1):
        ws.cell(i, 1).font = Font(bold=True, size=10, name="맑은 고딕")
        for j in (2, 3):
            ws.cell(i, j).font = BF
        ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(1, 1).font = Font(bold=True, size=15, name="맑은 고딕",
                              color="16324F")

    # ── ②~⑥ ──────────────────────────────────────────────────────
    sheet("② 측점", ST)
    v = V.copy().drop(columns=["출처"])
    v = v.rename(columns={"출처명": "출처", "rD분": "rD_분", "rI분": "rI_분"})
    cols = ["측점", "연도", "일자", "출처", "세션", "D", "I", "F",
            "rD_분", "rI_분", "rF_nT", "D_산포", "I_산포", "F_산포",
            "D판정", "I판정", "F판정", "좌표신뢰", "일자정밀도",
            "참방위각", "원본파일", "원본시트"]
    sheet("③ 방문 시계열", v[[c for c in cols if c in v.columns]]
          .sort_values(["측점", "연도", "출처"]),
          widths={"원본파일": 46})
    s2 = S.drop(columns=["출처"]).rename(
        columns={"출처명": "출처", "rD분": "rD_분", "rI분": "rI_분"})
    scols = ["측점", "연도", "일자", "출처", "회차", "D", "I", "F",
             "rD_분", "rI_분", "rF_nT", "D_판정", "I_판정", "F_판정",
             "D_원본", "D부호교정", "일자정밀도", "좌표신뢰", "참방위각",
             "위도", "경도", "원본파일", "원본시트"]
    sheet("④ 세션 원자료", s2[[c for c in scols if c in s2.columns]]
          .sort_values(["측점", "연도", "출처", "회차"]),
          widths={"원본파일": 46})
    sheet("⑤ 확인 필요", F, sevcol="심각도",
          widths={"왜": 54, "내용": 28, "원본파일": 46, "상세": 26})
    SR = pd.DataFrame(P["원본"])
    SR["있음"] = SR["있음"].map({True: "있음", False: "없음"})
    sheet("⑥ 원본 자료 위치",
          SR[["출처", "파일", "있음", "측점", "행", "최초", "최종",
              "폴더", "측점목록"]],
          widths={"폴더": 60, "측점목록": 40, "파일": 34})

    # ── ⑦ 잔차 그래프 — 측점별 편각 잔차를 연도축에 ──────────────
    #
    # ⚠️ 엑셀 차트는 **연속한 한 범위**를 읽는다. 측점이 행으로 섞여 있으면
    #    한 선으로 이어져 버리므로, 측점을 **열로 돌려** 깔아 둔다.
    piv = (v[v["rD_분"].notna()]
           .pivot_table(index="연도", columns="측점", values="rD_분",
                        aggfunc="median").sort_index())
    ws = wb.create_sheet("⑦ 잔차 그래프")
    ws.append(["연도"] + list(piv.columns))
    for c in range(1, piv.shape[1] + 2):
        ws.cell(1, c).fill, ws.cell(1, c).font = HEAD, HF
    for y, row in piv.iterrows():
        ws.append([int(y)] + [jnum(x, 2) for x in row.tolist()])
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 8
    # ⚠️ **30계열 선 그래프는 읽히지 않는다.** 색이 서로 구분되지 않고
    #    범례가 그림만큼 커진다. 한 장에 전부 그리는 대신 둘로 나눈다 —
    #    ① 전 측점을 한눈에 보는 **막대**(잔차 폭 순) ② 가장 흔들린
    #    **8측점만** 선으로. 나머지 조합은 왼쪽 피벗표로 직접 그린다.
    top = (piv.std().sort_values(ascending=False).head(8).index.tolist())
    cols = {c: i + 2 for i, c in enumerate(piv.columns)}

    ch = LineChart()
    ch.title = f"편각 IGRF 잔차 — 가장 흔들린 {len(top)}측점 (분)"
    ch.y_axis.title = "편각 잔차 (분)"
    ch.x_axis.title = "연도"
    ch.height, ch.width = 10, 26
    for c_ in top:
        ch.add_data(Reference(ws, min_col=cols[c_], min_row=1,
                              max_row=len(piv) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(piv) + 1))
    for s_ in ch.series:
        s_.smooth = False
    ws.add_chart(ch, f"A{len(piv) + 4}")

    # 전 측점 한눈에 — 잔차 폭 막대. 0 에 가까울수록 같은 자리를 다시 재도
    # 같은 값이 나왔다는 뜻이다.
    bw = ST[["측점", "rD_폭"]].dropna().sort_values("rD_폭", ascending=False)
    c0 = len(piv.columns) + 3
    ws.cell(1, c0, "측점").fill = HEAD
    ws.cell(1, c0).font = HF
    ws.cell(1, c0 + 1, "편각 잔차 폭(분)").fill = HEAD
    ws.cell(1, c0 + 1).font = HF
    for i, (_, r) in enumerate(bw.iterrows(), start=2):
        ws.cell(i, c0, r["측점"])
        ws.cell(i, c0 + 1, jnum(r["rD_폭"], 1))
    ws.column_dimensions[get_column_letter(c0)].width = 10
    ws.column_dimensions[get_column_letter(c0 + 1)].width = 16
    bc = BarChart()
    bc.type = "bar"
    bc.title = "측점별 편각 잔차 폭 (분) — 0 에 가까울수록 재방문이 일치한다"
    bc.height, bc.width = 14, 20
    bc.legend = None
    bc.add_data(Reference(ws, min_col=c0 + 1, min_row=1, max_row=len(bw) + 1),
                titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=c0, min_row=2, max_row=len(bw) + 1))
    ws.add_chart(bc, f"A{len(piv) + 26}")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main():
    P, (ST, V, S, F) = payload()
    OUT.mkdir(parents=True, exist_ok=True)

    h = OUT / "기존측정점_자료점검.html"
    h.write_text(build_html(P), encoding="utf-8")
    x = build_excel(ST, V, S, F, P, OUT / "기존측정점_자료점검.xlsx")

    print(f"■ 측점 {len(ST)} · 방문 {len(V)} · 세션 {len(S)}"
          f" · 플래그 {len(F)} · 원본 파일 {len(P['원본'])}")
    print(f"  심각도 — " + " · ".join(
        f"{k} {int((F['심각도'] == k).sum())}"
        for k in ("확인필요", "주의", "참고")))
    print(f"\n[저장] {h}  ({h.stat().st_size / 1024:.0f} KB)")
    print(f"[저장] {x}  ({x.stat().st_size / 1024:.0f} KB)")
    print("\n⚠ local/ 은 .gitignore 로 막혀 있다 — 커밋되지 않는다")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
